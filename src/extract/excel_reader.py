"""
Excel reader module - robustly reads and extracts data from shortages or warehouse sheets.
"""

from pathlib import Path
import re
from openpyxl import load_workbook


def read_excel(filepath: Path) -> list[dict]:
    filepath = Path(filepath)
    if not filepath.exists():
        raise FileNotFoundError(f"File not found: {filepath}")

    print(f"  -> Reading Excel file: {filepath.name}")

    try:
        wb = load_workbook(filename=filepath, data_only=True, read_only=True)
    except Exception as e:
        raise RuntimeError(f"Failed to read Excel file {filepath.name}: {e}")

    try:
        sheet = None
        for name in wb.sheetnames:
            ws = wb[name]
            if ws.max_row > 1 and ws.max_column > 0:
                sheet = ws
                break

        if not sheet:
            sheet = wb.active

        if not sheet or sheet.max_row <= 1:
            print("  [!] The Excel sheet appears to be empty.")
            return []

        print(f"  -> Active sheet: {sheet.title}")

        header_row_idx = None
        headers = []

        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            valid_cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if len(valid_cells) >= 2:
                header_row_idx = row_idx
                for idx, cell in enumerate(row):
                    val = str(cell).strip() if cell is not None else f"Column_{idx+1}"
                    if not val:
                        val = f"Column_{idx+1}"
                    headers.append(val)
                break

        if not header_row_idx:
            print("  [!] Could not detect a valid header row.")
            return []

        print(f"  -> Header row: {header_row_idx} -> {headers}")

        data = []
        for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            
            row_dict = {}
            for col_idx, header in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else None
                if val is not None and isinstance(val, str):
                    val = val.strip()
                row_dict[header] = val
                
            data.append(row_dict)

        print(f"  -> Extracted {len(data)} data rows")
        return data

    finally:
        wb.close()


def extract_item_names(rows: list[dict], name_column: str = None, source_file: str = "") -> list[dict]:
    if not rows:
        return []

    headers = list(rows[0].keys())

    if not name_column:
        best_col = ""
        best_score = -1

        keywords = ["صنف", "اسم", "دواء", "البيان", "item", "name", "product", "drug"]
        arabic_pattern = re.compile(r"[\u0600-\u06FF]")

        for header in headers:
            score = 0
            header_lower = str(header).lower()
            
            if any(kw in header_lower for kw in keywords):
                score += 50
                
            sample_vals = [str(r.get(header, "")) for r in rows[:10] if r.get(header)]
            if sample_vals:
                avg_len = sum(len(v) for v in sample_vals) / len(sample_vals)
                score += min(avg_len, 30)
                
                arabic_count = sum(1 for v in sample_vals if arabic_pattern.search(v))
                if arabic_count > 0:
                    score += (arabic_count / len(sample_vals)) * 20
            
            if score > best_score:
                best_score = score
                best_col = header
                
        if best_col:
            name_column = best_col
        else:
            name_column = headers[0]
            
        print(f"  -> Auto-detected drug name column: {name_column}")
    else:
        if name_column not in headers:
            raise ValueError(f"Column '{name_column}' not found. Available columns: {headers}")

    items = []
    for i, row in enumerate(rows, start=1):
        val = row.get(name_column)
        if val:
            items.append({
                "item_name_raw": str(val),
                "source_row": i,
                "source_file": source_file,
            })

    print(f"  -> Extracted {len(items)} items from column '{name_column}'")
    return items
