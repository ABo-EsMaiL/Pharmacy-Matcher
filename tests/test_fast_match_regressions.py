"""Historical fixtures under the 2026-09-28 presence policy and hostile provider decisions."""
import contextlib
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from src.config import Config
from src.fast_match.coordinator import (
    FastCoordinator, classify_pair, get_strength_pattern, validate_llm_decision,
)
from src.fast_match.search import FastCandidateFinder, extract_brand_tokens, extract_form_group
from src.output.excel_writer import write_results_excel
from src.match.text_match import normalize_text


class ParsingTests(unittest.TestCase):
    def test_proc017_decisions(self):
        pairs = [
            ('دولفين 50مجم لبوس 2شريط', 'دولفين 25لبوس 2شريط/باكو90', 'review'),
            ('دولفين 12.5مجم لبوس 2شريط س ق', 'دولفين 25لبوس 2شريط/باكو90', 'ambiguous'),
            ('لبن هيرو بيبي 2', 'هيرو بيبي ١ عادى', 'review'),
            ('فلافيسيف شراب30مل', 'فلافيسيف 60 مل شراب باكت 48', 'match'),
            ('Flacysiv syrup 30ml', 'Flacysiv syrup 60ml', 'match'),
        ]
        for requested, candidate, expected in pairs:
            with self.subTest(requested=requested):
                self.assertEqual(classify_pair(requested, candidate)[0], expected)

    def test_digits_and_decimals(self):
        for value in ['12.5mg', '۱۲,۵mg', '١٢٫٥مجم', '١٢،٥ مجم', '１２，５ｍｇ']:
            with self.subTest(value=value):
                self.assertEqual(get_strength_pattern(value), '12.5')
        self.assertEqual(get_strength_pattern('كولميديتين ٠,٥ مجم ۱۰ شريط'), '0.5')
        self.assertEqual(get_strength_pattern('دولفين ۲۵لبوس'), '25')

    def test_pack_volume_price_are_not_strength(self):
        for value in ['باكت25', 'باكو 50', '20 قرص', '6امبول', '2شريط',
                      'شراب30مل', 'شراب60ml', 'سعر 25', '25 جنيه',
                      'س.ج 25', 'جل 30 جم', 'Hero Baby 1', 'هيرو بيبي ٢']:
            with self.subTest(value=value):
                self.assertIsNone(get_strength_pattern(value))
        self.assertEqual(classify_pair('ليفابيون 6امبول', 'ليفابيون أمبول باكت 25')[0], 'match')

    def test_strength_equivalence_and_concentration(self):
        for left, right in [('1gm', '1000mg'), ('0,5 جم', '500 مجم'),
                            ('250mcg', '0.25mg'), ('50mg/5ml', '10mg/ml'),
                            ('875/125 mg', '875/125مجم')]:
            with self.subTest(left=left):
                self.assertEqual(get_strength_pattern(left), get_strength_pattern(right))
        self.assertEqual(classify_pair('دواء 50mg/5ml syrup', 'دواء 25mg/5ml syrup')[0], 'review')
        self.assertEqual(classify_pair('دواء 1% كريم', 'دواء 2% كريم')[0], 'review')

    def test_product_variants_survive(self):
        for left, right in [('Hero Baby 1', 'Hero Baby 2'), ('هيرو بيبي ١', 'هيرو بيبي ۲'),
                            ('Vitamin B12', 'Vitamin B6'), ('Vitamin B12', 'Vitamin D12')]:
            with self.subTest(left=left):
                self.assertNotEqual(extract_brand_tokens(left), extract_brand_tokens(right))
        self.assertEqual(classify_pair('هيرو بيبي ۲', 'لبن هيرو بيبي ٢ عادى')[0], 'match')

    def test_forms_and_packaging(self):
        self.assertIsNone(extract_form_group('2شريط'))
        self.assertEqual(extract_form_group('دولفين 25لبوس 2شريط'), 'suppository')
        self.assertEqual(extract_form_group('شراب30مل'), 'syrup')
        self.assertEqual(extract_form_group('شراب 30مل'), 'syrup')
        self.assertEqual(classify_pair('دواء 5mg كريم', 'دواء 5mg مرهم')[0], 'review')
        self.assertEqual(classify_pair('دواء 5mg اقراص', 'دواء 5mg كبسول')[0], 'match')

    def test_retrieval_preserves_form_difference(self):
        finder = FastCandidateFinder()
        finder.fit([{'id': 'W1', 'name': 'فاركوتيليام كبسول/باكت60'}])
        self.assertEqual(finder.search('فاركوتيليام شراب')[0]['id'], 'W1')
        self.assertEqual(classify_pair('فاركوتيليام شراب', 'فاركوتيليام كبسول/باكت60')[0], 'review')


class GuardTests(unittest.TestCase):
    def check(self, requested, candidate, status='match'):
        case = {'case_id': 'Q1', 'requested_name': requested, 'candidates': [{'id': 'W1', 'name': candidate}]}
        decision = {'case_id': 'Q1', 'status': status, 'item_id': 'W1'}
        return validate_llm_decision(case, decision)

    def test_llm_cannot_override_strength_or_combination(self):
        for status in ['match', 'review']:
            self.assertEqual(self.check('دولفين 50mg لبوس', 'دولفين 25لبوس', status)['status'], 'review')
            self.assertEqual(self.check('اتاكند بلس 8mg', 'اتاكند 8mg', status)['status'], 'review')

    def test_llm_adjudicates_plausible_ocr(self):
        result = self.check('بروفين 50mg', 'برونفين 50mg', 'match')
        self.assertEqual(result['status'], 'match')
        self.assertEqual(self.check('بروفين 50mg كريم', 'برونفين 50mg مرهم')['status'], 'review')
        self.assertEqual(self.check('Panadol 500mg', 'Cetal 500mg')['status'], 'no_match')

    def test_missing_details_are_adjudicated_not_auto_matched(self):
        requested, candidate = 'بروفين 400 مجم 20 قرص', 'بروفين'
        self.assertEqual(classify_pair(requested, candidate)[0], 'ambiguous')
        for status in ['match', 'no_match']:
            self.assertEqual(self.check(requested, candidate, status)['status'], 'review' if status == 'no_match' else status)
        self.assertEqual(classify_pair('بروفين 400mg', 'بروفين 600mg')[0], 'review')
        self.assertEqual(self.check('بروفين 400mg شراب', 'بروفين شراب', 'review')['status'], 'review')

    def test_python_owns_form_decision(self):
        self.assertEqual(self.check('دواء 5mg كريم', 'دواء 5mg مرهم')['status'], 'review')
        self.assertEqual(self.check('فلافيسيف شراب30مل', 'فلافيسيف شراب60مل', 'review')['status'], 'match')

    def test_outside_or_invalid_ids_and_status_fail(self):
        case = {'case_id': 'Q1', 'requested_name': 'دواء', 'candidates': [{'id': 'W1', 'name': 'دواء'}]}
        for item_id in ['W2', '', None, ['W1']]:
            with self.subTest(item_id=item_id), self.assertRaises(ValueError):
                validate_llm_decision(case, {'case_id': 'Q1', 'status': 'match', 'item_id': item_id})
        with self.assertRaises(ValueError):
            validate_llm_decision(case, {'case_id': 'Q1', 'status': 'invented', 'item_id': 'W1'})


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.engine = FastCoordinator(Config('', '', local_api_url='http://127.0.0.1:8001/v1/chat/completions', local_model='gemini-3.8-flash'))
        self.silence = contextlib.redirect_stdout(io.StringIO())
        self.silence.__enter__()
        self.addCleanup(self.silence.__exit__, None, None, None)
        self.sleep = patch('src.fast_match.coordinator.time.sleep')
        self.sleep.start()
        self.addCleanup(self.sleep.stop)
        self.network = patch('src.fast_match.coordinator.http_requests.post', side_effect=AssertionError('Network forbidden in regressions'))
        self.network.start()
        self.addCleanup(self.network.stop)

    def test_rejected_candidates_removed_sixth_preserved(self):
        candidates = [{'id': f'W{i}', 'name': 'بروفين 25mg', 'score': .8} for i in range(5)]
        candidates.append({'id': 'W5', 'name': 'برونفين 50mg', 'score': .7})
        def respond(inputs):
            self.assertEqual([(c['id'], c['name']) for c in inputs['cases'][0]['candidates']], [('W5', 'برونفين 50mg')])
            return [{'case_id': 'Q00000', 'status': 'match', 'item_id': 'W5'}]
        with patch('src.fast_match.coordinator.FastCandidateFinder.search', return_value=candidates), patch.object(self.engine, '_call_local_api', side_effect=respond):
            result = self.engine._match_warehouse([{'item_name_raw': 'بروفين 50mg'}], [{'item_name_raw': 'برونفين 50mg'}], 'fixture')
        self.assertEqual(len(result['matched']), 1)
        self.assertEqual(result['matched'][0]['warehouse_item'], 'برونفين 50mg')
        self.assertEqual(result['needs_review'], [])
        self.assertNotIn('identity_audit', result)

    def test_provider_failure_never_becomes_review(self):
        for message in ['HTTP 429: Rate Limited', 'HARD_CEILING_TIMEOUT', 'PROVIDER_RESPONSE_INVALID']:
            with self.subTest(message=message), patch.object(self.engine, '_call_local_api', side_effect=RuntimeError(message)), self.assertRaisesRegex(RuntimeError, 'Provider failed'):
                self.engine._match_warehouse([{'item_name_raw': 'بروفين 50mg'}], [{'item_name_raw': 'برونفين 50mg'}], 'fixture')

    def test_deterministic_path_no_provider(self):
        with patch.object(self.engine, '_call_local_api', side_effect=AssertionError('Unexpected LLM')):
            result = self.engine.process([{'item_name_raw': 'فلافيسيف شراب30مل'}, {'item_name_raw': 'دولفين 50مجم لبوس'}], {'fixture': [{'item_name_raw': 'فلافيسيف شراب60مل'}, {'item_name_raw': 'دولفين 25لبوس'}]})
        self.assertEqual(result['summary']['matched_count'], 1)
        self.assertEqual(result['summary']['review_count'], 1)
        self.assertEqual(result['summary']['not_found_count'], 0)

    def test_excel_has_only_business_outcomes(self):
        with patch.object(self.engine, '_call_local_api', return_value=[{'case_id': 'Q00000', 'status': 'match', 'item_id': 'W-00000'}]):
            result = self.engine.process([{'item_name_raw': 'بروفين 50mg'}], {'fixture': [{'item_name_raw': 'برونفين 50mg'}]})
        self.assertNotIn('brand_identity_unverified_count', result['summary'])
        self.assertEqual(result['not_found'], [])
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'regression.xlsx'
            write_results_excel(result, path)
            book = load_workbook(path, read_only=True)
            try:
                self.assertNotIn('تدقيق هوية البراند', book.sheetnames)
                self.assertEqual(book['يحتاج مراجعة'].max_row, 1)
                self.assertEqual(book['fixture'].max_row, 2)
            finally:
                book.close()

    def test_llm_no_match_is_final_business_outcome(self):
        with patch.object(self.engine, '_call_local_api', return_value=[{'case_id': 'Q00000', 'status': 'no_match', 'item_id': '', 'identity_same': False}]):
            result = self.engine.process([{'item_name_raw': 'بروفين 50mg'}], {'fixture': [{'item_name_raw': 'برونفين 50mg'}]})
        self.assertNotIn('brand_identity_unverified_count', result['summary'])
        self.assertEqual(result['summary']['matched_count'], 0)
        self.assertEqual(result['summary']['review_count'], 0)
        self.assertEqual(result['not_found'][0]['name'], 'بروفين 50mg')

    def test_missing_fields_reach_llm(self):
        requested = 'بروفين 400 مجم 20 قرص'
        with patch.object(self.engine, '_call_local_api', return_value=[{'case_id': 'Q00000', 'status': 'match', 'item_id': 'W-00000'}]) as provider:
            result = self.engine.process([{'item_name_raw': requested}], {'fixture': [{'item_name_raw': 'بروفين'}]})
        self.assertEqual(provider.call_count, 1)
        self.assertEqual(provider.call_args.args[0]['cases'][0]['requested_name'], requested)
        self.assertEqual(result['summary']['matched_count'], 1)

    def test_only_ambiguous_candidates_sent_with_deterministic_review_fallback(self):
        candidates = [
            {'id': 'form', 'name': 'zorvanel 5mg cream', 'score': .9},
            {'id': 'ocr', 'name': 'zorvaneli 5mg syrup', 'score': .8},
        ]
        def respond(inputs):
            self.assertEqual([c['id'] for c in inputs['cases'][0]['candidates']], ['ocr'])
            return [{'case_id': 'Q00000', 'status': 'no_match', 'item_id': '', 'identity_same': False}]
        with patch('src.fast_match.coordinator.FastCandidateFinder.search', return_value=candidates), patch.object(self.engine, '_call_local_api', side_effect=respond):
            result = self.engine._match_warehouse([{'item_name_raw': 'zorvanel 5mg syrup'}], [{'item_name_raw': c['name']} for c in candidates], 'fixture')
        self.assertEqual(result['matched'], [])
        self.assertEqual(result['needs_review'][0]['warehouse_item'], 'zorvanel 5mg cream')

    def test_uncertain_exact_text_does_not_disappear_before_llm(self):
        name = 'zorvanel syrup tablets'
        with patch.object(self.engine, '_call_local_api', return_value=[{'case_id': 'Q00000', 'status': 'match', 'item_id': 'W-00000'}]) as provider:
            result = self.engine.process([{'item_name_raw': name}], {'fixture': [{'item_name_raw': name}]})
        provider.assert_called_once()
        self.assertEqual(result['summary']['matched_count'], 1)

    def test_no_name_evidence_never_reaches_llm(self):
        with patch.object(self.engine, '_call_local_api', side_effect=AssertionError('Garbage candidate')) as provider:
            result = self.engine.process([{'item_name_raw': 'كارديكسين امبول'}], {'fixture': [{'item_name_raw': 'جرامايسين كريم'}, {'item_name_raw': 'جاراميسين 80 مج امبول/باكت 400'}]})
        provider.assert_not_called()
        self.assertEqual(result['summary']['matched_count'], 0)
        self.assertEqual(result['summary']['not_found_count'], 1)


class FullNameEvidenceTests(unittest.TestCase):
    def test_unicode_presentation_forms_and_persian_letters(self):
        requested = 'كارديکسين امبول'
        source = 'ﻛﺎردﯾﻛﺳﯾن اﻣﺑول'
        self.assertEqual(normalize_text(requested), normalize_text(source))
        self.assertEqual(classify_pair(requested, source)[0], 'match')
        finder = FastCandidateFinder()
        finder.fit([{'id': 'same', 'name': source}, {'id': 'wrong', 'name': 'جرامايسين كريم'}])
        self.assertEqual([c['id'] for c in finder.search(requested)], ['same'])
        self.assertEqual(classify_pair('Ｃｏ zorvanel 5mg', 'Co zorvanel 5mg')[0], 'match')

    def test_generic_ocr_mutations_retain_recall(self):
        # Synthetic names exercise the algorithm independently of drug fixtures.
        for name in ['zorvanel', 'lumetaris', 'navorelin', 'تولمارين']:
            mutations = [name[1:], name[:3] + name[4:], name[:3] + 'x' + name[3:],
                         name[:3] + ' ' + name[3:]]
            for mutated in mutations:
                with self.subTest(name=name, mutated=mutated):
                    finder = FastCandidateFinder()
                    finder.fit([{'id': 'ocr', 'name': mutated + ' 5mg syrup'}])
                    self.assertEqual([c['id'] for c in finder.search(name + ' 5mg syrup')], ['ocr'])
                    self.assertEqual(classify_pair(name + ' 5mg syrup', mutated + ' 5mg syrup')[0],
                                     'match' if mutated.replace(' ', '') == name else 'ambiguous')

    def test_shared_suffix_prefix_and_form_do_not_admit_garbage(self):
        cases = [('tavorelin', 'xuzarelin'), ('tavorelin', 'tavo'),
                 ('alpha medication', 'alpha unrelated'), ('كارديكسين', 'جرامايسين'),
                 ('كارديكسين', 'جراميسين'), ('كارديكسين', 'جاراميسين')]
        for left, right in cases:
            with self.subTest(left=left, right=right):
                finder = FastCandidateFinder()
                finder.fit([{'id': 'bad', 'name': right + ' 50mg syrup 30ml box20'}])
                self.assertEqual(finder.search(left + ' 50mg syrup 30ml box20'), [])

    def test_ranking_survives_packaging_and_distractors(self):
        finder = FastCandidateFinder()
        finder.fit([{'id': str(i), 'name': f'irrelevant{i} syrup 30ml box20'} for i in range(30)] +
                   [{'id': 'real', 'name': 'z orvanel 50mg syrup 60ml bag40'}])
        self.assertEqual(finder.search('zorvanel 50mg syrup 30ml box20')[0]['id'], 'real')

    def test_uncertain_forms_are_ambiguous(self):
        self.assertEqual(classify_pair('zorvanel syrup tablets', 'zorvanel syrup')[0], 'ambiguous')

    def test_arabic_single_letter_product_parts_are_preserved(self):
        self.assertNotEqual(extract_brand_tokens('فيتامين ب12'), extract_brand_tokens('فيتامين د12'))


if __name__ == '__main__':
    unittest.main()
