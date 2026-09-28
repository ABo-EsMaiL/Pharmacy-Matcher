import json
import os
import shutil
import subprocess
import sys
import threading
from pathlib import Path
import openpyxl

# Add project root to path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from desktop_app.database import HistoryDB
from src.config import gateway_api_key
from desktop_app.server_manager import MSEMAXServerManager

class LogRedirector:
    def __init__(self, original_stream, add_log_func):
        self.original_stream = original_stream
        self.add_log_func = add_log_func
        self.buffer = ""
        
    def write(self, message):
        self.original_stream.write(message)
        self.buffer += message
        while "\n" in self.buffer:
            line, self.buffer = self.buffer.split("\n", 1)
            self.add_log_func(line)
            
    def flush(self):
        self.original_stream.flush()

    def isatty(self):
        return False

class API:
    def __init__(self, window=None):
        self._window = window
        self._db = HistoryDB()
        
        # Clean up any processes stuck in 'processing' state from a previous crash/close
        try:
            self._db.mark_stuck_processes_as_failed()
        except:
            pass
            
        self._settings_file = PROJECT_ROOT / 'data' / 'settings.json'
        self._logs = []
        self._server_logs = []
        self._dropped_files = []
        
        settings = self.get_settings()
        self._server_manager = MSEMAXServerManager(settings, server_logger_callback=self.add_server_log)
        
        # Redirect stdout and stderr so that prints show up in the pipeline logs
        sys.stdout = LogRedirector(sys.stdout, self.add_log)
        sys.stderr = LogRedirector(sys.stderr, self.add_log)
        
        self._worker_thread = None
        self._worker_lock = threading.Lock()
        self._process_queue = []

    def set_window(self, window):
        """Set pywebview window reference safely without exposing it to JS reflection."""
        self._window = window
        
    def add_log(self, message: str):
        self._logs.append(message)
        if len(self._logs) > 1000:
            self._logs.pop(0)
        
        if self._window:
            try:
                safe_msg = message.replace('\\', '\\\\').replace("'", "\\'").replace('\n', '\\n').replace('\r', '')
                self._window.evaluate_js(f"window.dispatchEvent(new CustomEvent('new_log', {{detail: '{safe_msg}'}}));")
            except:
                pass

    def get_logs(self) -> str:
        return json.dumps(self._logs)

    def add_server_log(self, message: str):
        self._server_logs.append(message)
        if len(self._server_logs) > 1000:
            self._server_logs.pop(0)
        
        if self._window:
            try:
                safe_msg = message.replace('\\', '\\\\').replace("'", "\\'").replace('\n', '\\n').replace('\r', '')
                self._window.evaluate_js(f"window.dispatchEvent(new CustomEvent('new_server_log', {{detail: '{safe_msg}'}}));")
            except:
                pass

    def get_server_logs(self) -> str:
        return json.dumps(self._server_logs)

    def clear_server_logs(self) -> dict:
        self._server_logs.clear()
        return {"success": True}

    def save_dropped_file(self, filename: str, base64_data: str) -> str:
        import base64
        import tempfile
        from pathlib import Path
        try:
            data = base64.b64decode(base64_data)
            temp_dir = Path(tempfile.gettempdir()) / "pharmacy_matcher_drops"
            temp_dir.mkdir(parents=True, exist_ok=True)
            
            safe_name = "".join([c for c in filename if c.isalpha() or c.isdigit() or c in (' ', '.', '-', '_')]).rstrip()
            if not safe_name:
                safe_name = "dropped_file.tmp"
                
            file_path = temp_dir / safe_name
            with open(file_path, "wb") as f:
                f.write(data)
                
            return str(file_path)
        except Exception as e:
            self.add_log(f"Error saving dropped file: {e}")
            return ""

    def exit_app(self):
        try:
            self.on_closing()
        except:
            pass
        # os._exit guarantees immediate shutdown without blocking background threads like uvicorn
        os._exit(0)
            
    def get_settings(self) -> dict:
        if self._settings_file.exists():
            try:
                with open(self._settings_file, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    # Normalize local_api_url if missing /chat/completions
                    if "local_api_url" in data:
                        raw = str(data["local_api_url"]).strip().rstrip("/")
                        if not raw.endswith("/chat/completions"):
                            if raw.endswith("/v1"):
                                data["local_api_url"] = f"{raw}/chat/completions"
                            else:
                                data["local_api_url"] = f"{raw}/v1/chat/completions"
                    return data
            except:
                pass
        return {
            "unstructured_api_key": "",
            "local_api_url": "http://127.0.0.1:8001/v1/chat/completions",
            "local_api_key": gateway_api_key(),
            "local_model": "gemini-3.8-flash",
            "vision_api_url": "http://127.0.0.1:8001/v1/chat/completions",
            "vision_api_key": gateway_api_key(),
            "vision_model": "gemini-3.8-flash",
            "active_server": "gaistudio",
            "msemax_dir": r"D:\AI_Engineer\MSEMAX",
            "msemax_chatgpt_dir": r"D:\AI_Engineer\MSEMAX",
            "msemax_gaistudio_dir": r"D:\AI_Engineer\MSEMAX-GAIStudio-vision",
            "unstructured_profile": "balanced"
        }

    def save_settings(self, settings_json: str) -> dict:
        try:
            settings = json.loads(settings_json)
            # Ensure URL has /chat/completions
            if "local_api_url" in settings:
                raw = str(settings["local_api_url"]).strip().rstrip("/")
                if not raw.endswith("/chat/completions"):
                    if raw.endswith("/v1"):
                        settings["local_api_url"] = f"{raw}/chat/completions"
                    else:
                        settings["local_api_url"] = f"{raw}/v1/chat/completions"

            self._settings_file.parent.mkdir(parents=True, exist_ok=True)
            with open(self._settings_file, 'w', encoding='utf-8') as f:
                json.dump(settings, f, ensure_ascii=False, indent=2)
            
            # Refresh server manager with latest settings
            self._server_manager = MSEMAXServerManager(settings, server_logger_callback=self.add_server_log)
            return {"success": True}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def get_servers_info(self) -> str:
        """Returns JSON of all servers status and configuration."""
        settings = self.get_settings()
        active_key = settings.get("active_server", "gaistudio")
        statuses = self._server_manager.get_all_status(active_key=active_key)
        return json.dumps({
            "servers": statuses,
            "active_server": active_key,
            "current_url": settings.get("local_api_url", ""),
            "current_model": settings.get("local_model", "")
        }, ensure_ascii=False)

    def start_server_instance(self, server_key: str) -> dict:
        """Starts a specific server by key ('chatgpt' or 'gaistudio')."""
        return self._server_manager.start(server_key)

    def stop_server_instance(self, server_key: str) -> dict:
        """Stops a specific server by key."""
        return self._server_manager.stop(server_key)

    def test_server_instance(self, server_key: str) -> dict:
        """Sends a test ping to the specified server and measures latency."""
        return self._server_manager.test(server_key)

    def run_server_session(self, server_key: str) -> dict:
        """Launches get_session.py for the specified server."""
        return self._server_manager.run_get_session(server_key)

    def free_server_ports(self) -> dict:
        """Force kills any process occupying ports 8000/8001."""
        self.add_server_log("[SYSTEM] Force freeing ports 8000 and 8001...")
        res = self._server_manager.force_free_ports()
        self.add_server_log(f"[SYSTEM] Free ports operation completed: {res}")
        return res

    def set_active_ai_server(self, server_key: str) -> dict:
        """
        1-Click Active Switch:
        Switches which server will receive matching batches.
        Updates settings.json (local_api_url, local_model, active_server) automatically.
        """
        try:
            settings = self.get_settings()
            srv = self._server_manager.get_server(server_key)
            if not srv:
                return {"success": False, "error": f"Server {server_key} not found"}

            settings["active_server"] = server_key
            settings["local_api_url"] = srv.url
            settings["local_model"] = srv.default_model

            self.save_settings(json.dumps(settings))
            self.add_log(f"[SYSTEM] Switched active AI server to: {srv.name} (Port: {srv.port}, Model: {srv.default_model})")

            # Auto-start target server if currently stopped
            if not srv.is_running():
                self.add_log(f"[SYSTEM] Server {srv.name} is not running. Starting it automatically...")
                start_res = srv.start()
                if start_res.get("success"):
                    self.add_log(f"[SYSTEM] Server {srv.name} started successfully (PID: {start_res.get('pid')}).")
                else:
                    self.add_log(f"[WARNING] Could not auto-start {srv.name}: {start_res.get('message')}")

            return {"success": True, "active_server": server_key, "url": srv.url, "model": srv.default_model}
        except Exception as e:
            return {"success": False, "error": str(e)}
    def update_server_dir(self, server_key: str, new_dir: str) -> dict:
        """Updates directory path for a server instance and saves to settings."""
        settings = self.get_settings()
        if server_key == "gaistudio":
            settings["msemax_gaistudio_dir"] = new_dir
            settings["msemax_dir"] = new_dir
        elif server_key == "chatgpt":
            settings["msemax_chatgpt_dir"] = new_dir
        self.save_settings(json.dumps(settings))
        self._server_manager = MSEMAXServerManager(settings, server_logger_callback=self.add_server_log)
        self.add_log(f"[SETTINGS] Updated {server_key} folder to: {new_dir}")
        return {"success": True, "dir": new_dir}

    # Legacy compatibility methods
    def start_server(self) -> dict:
        settings = self.get_settings()
        active_key = settings.get("active_server", "gaistudio")
        return self._server_manager.start(active_key)

    def stop_server(self) -> dict:
        settings = self.get_settings()
        active_key = settings.get("active_server", "gaistudio")
        return self._server_manager.stop(active_key)

    def get_server_status(self) -> dict:
        settings = self.get_settings()
        active_key = settings.get("active_server", "gaistudio")
        srv = self._server_manager.get_server(active_key)
        return srv.get_status() if srv else {"running": False}

    def run_get_session(self) -> dict:
        settings = self.get_settings()
        active_key = settings.get("active_server", "gaistudio")
        return self._server_manager.run_get_session(active_key)

    def select_files(self, file_type: str) -> str:
        """Opens a native file dialog. Returns JSON list of selected file paths."""
        import webview
        if not self._window:
            return "[]"
        
        if file_type == 'warehouse':
            file_types = ('Warehouse Files (*.pdf;*.xlsx;*.xls;*.json)', 'All Files (*.*)')
        else:
            file_types = ('Shortage Files (*.pdf;*.xlsx;*.xls)', 'All Files (*.*)')
        
        try:
            from webview import FileDialog
            dialog_type = FileDialog.OPEN
        except ImportError:
            import webview
            dialog_type = webview.OPEN_DIALOG
        
        result = self._window.create_file_dialog(
            dialog_type,
            allow_multiple=True,
            file_types=file_types
        )
        
        if result:
            return json.dumps([str(p) for p in result])
        return "[]"

    def select_directory(self) -> str:
        """Opens a native directory picker. Returns selected path."""
        import webview
        if not self._window:
            return ""
            
        try:
            from webview import FileDialog
            dialog_type = FileDialog.FOLDER
        except ImportError:
            import webview
            dialog_type = webview.FOLDER_DIALOG
            
        result = self._window.create_file_dialog(dialog_type)
        if result and len(result) > 0:
            return str(result[0])
        return ""

    def is_process_running(self) -> dict:
        running = bool(self._worker_thread and self._worker_thread.is_alive())
        return {"running": running, "queue_length": len(self._process_queue)}

    def start_process(self, shortage_paths_json: str, warehouse_paths_json: str) -> dict:
        try:
            shortage_paths = json.loads(shortage_paths_json)
            warehouse_paths = json.loads(warehouse_paths_json)
            
            shortage_filenames = [Path(p).name for p in shortage_paths]
            warehouse_filenames = [Path(p).name for p in warehouse_paths]
            
            process_id = self._db.create_process(shortage_filenames, warehouse_filenames)
            
            with self._worker_lock:
                if self._worker_thread and self._worker_thread.is_alive():
                    self._db.update_process(process_id, status='queued')
                    self._process_queue.append((process_id, shortage_paths, warehouse_paths))
                    self.add_log(f"[QUEUE] Process [{process_id}] added to queue (position #{len(self._process_queue)})")
                    return {"success": True, "process_id": process_id, "queued": True, "position": len(self._process_queue)}
                else:
                    self._worker_thread = threading.Thread(target=self._process_worker, args=(process_id, shortage_paths, warehouse_paths))
                    self._worker_thread.daemon = True
                    self._worker_thread.start()
                    return {"success": True, "process_id": process_id, "queued": False}
                    
        except Exception as e:
            return {"success": False, "error": str(e)}

    def restart_process(self, process_id: str) -> dict:
        with self._worker_lock:
            if self._worker_thread and self._worker_thread.is_alive():
                return {"success": False, "error": "توجد عملية مطابقة قيد التشغيل حالياً، يرجى الانتظار حتى تنتهي"}

            try:
                workspace = PROJECT_ROOT / 'data' / 'inputs' / process_id
                if not workspace.exists():
                    return {"success": False, "error": "ملفات العملية غير موجودة لإعادة التشغيل"}
                
                # get process details
                cursor = self._db.conn.cursor()
                cursor.execute('SELECT shortage_files, warehouse_files FROM processes WHERE id = ?', (process_id,))
                row = cursor.fetchone()
                if not row:
                    return {"success": False, "error": "العملية غير موجودة"}
                    
                shortage_filenames = json.loads(row['shortage_files'])
                warehouse_filenames = json.loads(row['warehouse_files'])
                
                shortage_paths = [str(workspace / name) for name in shortage_filenames]
                warehouse_paths = [str(workspace / name) for name in warehouse_filenames]
                
                # Reset status
                self._db.update_process_status(process_id, 'processing')
                
                self._worker_thread = threading.Thread(target=self._process_worker, args=(process_id, shortage_paths, warehouse_paths, True))
                self._worker_thread.daemon = True
                self._worker_thread.start()
                
                return {"success": True, "process_id": process_id}
            except Exception as e:
                return {"success": False, "error": str(e)}

    def _process_worker(self, process_id: str, shortage_paths: list, warehouse_paths: list, is_restart: bool = False):
        try:
            self.add_log(f"--- Starting new process [{process_id}] ---")
            proc_dir = self._db.get_process_dir(process_id)
            
            # Create subdirs
            shortages_dir = proc_dir / 'input_shortages'
            shortages_dir.mkdir(parents=True, exist_ok=True)
            warehouses_dir = proc_dir / 'input_warehouses'
            warehouses_dir.mkdir(parents=True, exist_ok=True)
            
            # Copy files to workspace if not restarting
            copied_shortages = []
            for p in shortage_paths:
                dest = shortages_dir / Path(p).name
                if not is_restart or not dest.exists():
                    try:
                        shutil.copy2(p, dest)
                    except shutil.SameFileError:
                        pass
                copied_shortages.append(dest)
                
            copied_warehouses = []
            for p in warehouse_paths:
                dest = warehouses_dir / Path(p).name
                if not is_restart or not dest.exists():
                    try:
                        shutil.copy2(p, dest)
                    except shutil.SameFileError:
                        pass
                copied_warehouses.append(dest)
                
            if self._window:
                self._window.evaluate_js(f"window.dispatchEvent(new CustomEvent('process_update', {{detail: {{process_id: '{process_id}', status: 'processing', message: 'جاري فحص وقراءة الملفات...'}}}}));")
                
            # Import actual project modules
            from src.config import load_config
            from src.extract.file_handler import read_input_files
            from src.fast_match.coordinator import FastCoordinator
            from src.output.excel_writer import write_results_excel
            
            # Load and update config with current settings
            settings = self.get_settings()
            config = load_config()
            config.unstructured_api_key = settings.get('unstructured_api_key', '') or getattr(config, 'unstructured_api_key', '')
            config.vision_api_url = settings.get('vision_api_url', '') or getattr(config, 'vision_api_url', 'http://127.0.0.1:8001/v1/chat/completions')
            config.vision_api_key = settings.get('vision_api_key', '') or getattr(config, 'vision_api_key', gateway_api_key())
            config.vision_model = settings.get('vision_model', '') or getattr(config, 'vision_model', 'gemini-3.8-flash')
            config.local_api_url = settings.get('local_api_url', '') or config.local_api_url
            config.local_api_key = settings.get('local_api_key', '') or config.local_api_key
            config.local_model = settings.get('local_model', '') or config.local_model
            
            self.add_log("[PROCESS] Config loaded (Vision Gateway: active).")
            if self._window:
                self._window.evaluate_js(f"window.dispatchEvent(new CustomEvent('process_update', {{detail: {{process_id: '{process_id}', status: 'processing', message: 'جاري استخراج الأصناف من الفواتير والملفات...'}}}}));")
            
            # Read input files (Pure Gemini Vision for PDF, Smart Excel & Markdown parsers)
            vision_url = config.vision_api_url
            vision_key = config.vision_api_key
            vision_model = config.vision_model
            profile = settings.get('unstructured_profile', 'balanced')
            
            self.add_log(f"[PROCESS] Reading Shortages files (Multi-Format Engine)...")
            shortage_data = read_input_files(
                copied_shortages,
                role="shortages",
                vision_api_url=vision_url,
                vision_api_key=vision_key,
                vision_model=vision_model,
                profile=profile
            )
            self.add_log(f"[PROCESS] Read {len(shortage_data)} shortage files.")
            
            self.add_log(f"[PROCESS] Reading Warehouse files (Multi-Format / Vision Engine)...")
            warehouse_data = read_input_files(
                copied_warehouses,
                role="warehouse",
                vision_api_url=vision_url,
                vision_api_key=vision_key,
                vision_model=vision_model,
                profile=profile
            )
            self.add_log(f"[PROCESS] Read {len(warehouse_data)} warehouse files.")
            
            # Flatten shortages
            all_shortages = []
            for file_items in shortage_data.values():
                all_shortages.extend(file_items)
            
            self.add_log(f"[PROCESS] Extracted total {len(all_shortages)} shortage items. Starting coordinator...")
            
            # Ensure local AI server is active & responsive before matching
            active_server_key = settings.get("active_server", "chatgpt" if "8000" in getattr(config, 'local_api_url', '') else "gaistudio")
            srv = self._server_manager.get_server(active_server_key)
            if srv and hasattr(config, 'local_api_url') and config.local_api_url:
                if not srv.is_running():
                    self.add_log(f"[PROCESS] Local AI server ({srv.name}) on port {srv.port} is not running. Auto-starting...")
                    start_res = srv.start()
                    if start_res.get("success"):
                        self.add_log(f"[PROCESS] Server spawned (PID: {start_res.get('pid')}). Waiting for port readiness...")
                        for _ in range(12):
                            time.sleep(0.5)
                            if srv.is_running():
                                break
                    else:
                        self.add_log(f"[PROCESS WARNING] Could not auto-start {srv.name}: {start_res.get('message')}")

            if self._window:
                self._window.evaluate_js(f"window.dispatchEvent(new CustomEvent('process_update', {{detail: {{process_id: '{process_id}', status: 'processing', message: 'جاري مطابقة {len(all_shortages)} صنف...'}}}}));")
            
            # Run matching
            coordinator = FastCoordinator(config)
            results = coordinator.process(all_shortages, warehouse_data)
            
            self.add_log(f"[PROCESS] Matching completed.")
            if self._window:
                self._window.evaluate_js(f"window.dispatchEvent(new CustomEvent('process_update', {{detail: {{process_id: '{process_id}', status: 'processing', message: 'جاري حفظ نتائج المطابقة وإعداد التقرير...'}}}}));")
                
            # Write results
            out_file = proc_dir / f"results_{process_id}.xlsx"
            write_results_excel(results, out_file)
            
            # Update DB
            summary = results.get("summary", {})
            self._db.update_process(
                process_id,
                status='reviewing' if summary.get('review_count', 0) > 0 else 'completed',
                total_shortages=summary.get('total_shortages', 0),
                matched_count=summary.get('matched_count', 0),
                not_found_count=summary.get('not_found_count', 0),
                review_count=summary.get('review_count', 0),
                output_file=str(out_file)
            )
            
            self.add_log(f"[PROCESS] Saved results to {out_file.name}")
            self.add_log(f"--- Process completed successfully [{process_id}] ---")
            
            # Send completion event back to UI
            if self._window:
                safe_stats = json.dumps(summary).replace('\\', '\\\\').replace('"', '\\"')
                self._window.evaluate_js(f"window.dispatchEvent(new CustomEvent('process_complete', {{detail: {{process_id: '{process_id}', success: true, stats: JSON.parse('{safe_stats}')}}}}));")
                
        except Exception as e:
            import traceback
            traceback.print_exc()
            self.add_log(f"[PROCESS ERROR] Exception occurred: {str(e)}")
            self._db.update_process(process_id, status='error')
            err_msg = str(e).replace("'", "\\'").replace("\n", " ").replace('"', '\\"')
            if self._window:
                self._window.evaluate_js(f"window.dispatchEvent(new CustomEvent('process_complete', {{detail: {{process_id: '{process_id}', success: false, error: \"{err_msg}\"}}}}));")
        finally:
            with self._worker_lock:
                if self._process_queue:
                    next_id, next_shortages, next_warehouses = self._process_queue.pop(0)
                    self.add_log(f"[QUEUE] Triggering next process from queue: [{next_id}]")
                    self._db.update_process(next_id, status='processing')
                    if self._window:
                        try:
                            self._window.evaluate_js(f"window.dispatchEvent(new CustomEvent('process_queued_started', {{detail: {{process_id: '{next_id}'}}}}));")
                        except Exception:
                            pass
                    self._worker_thread = threading.Thread(
                        target=self._process_worker,
                        args=(next_id, next_shortages, next_warehouses, False),
                        daemon=True
                    )
                    self._worker_thread.start()
                else:
                    self._worker_thread = None

    def get_process_result(self, process_id: str) -> dict:
        proc = self._db.get_process(process_id)
        if proc:
            proc['shortage_files'] = json.loads(proc['shortage_files'])
            proc['warehouse_files'] = json.loads(proc['warehouse_files'])
            return {"success": True, "process": proc}
        return {"success": False, "error": "Process not found"}

    def open_excel(self, process_id: str) -> dict:
        proc = self._db.get_process(process_id)
        if proc and proc.get('output_file') and os.path.exists(proc['output_file']):
            try:
                os.startfile(proc['output_file'])
                return {"success": True}
            except Exception as e:
                return {"success": False, "error": str(e)}
        return {"success": False, "error": "File not found"}

    def open_folder(self, process_id: str) -> dict:
        try:
            folder = self._db.get_process_dir(process_id)
            if folder.exists():
                os.startfile(str(folder))
                return {"success": True}
            return {"success": False, "error": "Folder not found"}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def open_file(self, filepath: str) -> dict:
        try:
            if os.path.exists(filepath):
                os.startfile(filepath)
                return {"success": True}
            return {"success": False, "error": "File not found"}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def apply_reviews(self, process_id: str) -> dict:
        proc = self._db.get_process(process_id)
        if not proc or not proc.get('output_file') or not os.path.exists(proc['output_file']):
            return {"success": False, "error": "ملف النتائج غير موجود"}
            
        try:
            wb = openpyxl.load_workbook(proc['output_file'])
            if 'يحتاج مراجعة' not in wb.sheetnames:
                return {"success": False, "error": "شيت المراجعة غير موجود"}
            
            review_ws = wb['يحتاج مراجعة']
            not_found_ws = wb['لم يُعثر عليه'] if 'لم يُعثر عليه' in wb.sheetnames else None
            
            matched_count = 0
            rejected_count = 0
            rows_to_delete = []
            
            # Columns in review sheet: القرار(A), اسم الصنف المطلوب(B), اسم المرشح(C), المخزن(D), السبب(E)
            for row_idx in range(2, review_ws.max_row + 1):
                decision = review_ws.cell(row=row_idx, column=1).value
                req_name = review_ws.cell(row=row_idx, column=2).value
                wh_candidate = review_ws.cell(row=row_idx, column=3).value
                wh_name = review_ws.cell(row=row_idx, column=4).value
                
                if decision == '✅ مطابق':
                    # Find the warehouse sheet and add the item there
                    safe_wh_name = wh_name or ""
                    for ch in r"[]:*?/\\":
                        safe_wh_name = safe_wh_name.replace(ch, "_")
                    safe_wh_name = safe_wh_name[:31]
                    
                    if safe_wh_name in wb.sheetnames:
                        ws_target = wb[safe_wh_name]
                        ws_target.append([req_name, wh_candidate, "تأكيد يدوي (مراجعة)"])
                    
                    matched_count += 1
                    rows_to_delete.append(row_idx)
                    
                elif decision == '❌ غير مطابق':
                    # Move to "لم يُعثر عليه"
                    if not_found_ws:
                        not_found_ws.append([req_name, "مرفوض يدوياً"])
                    
                    rejected_count += 1
                    rows_to_delete.append(row_idx)
                    
                # Empty decision = leave as is (user hasn't decided yet)
                    
            # Delete processed rows from bottom to top
            for row_idx in sorted(rows_to_delete, reverse=True):
                review_ws.delete_rows(row_idx)
                
            # Recalculate totals across all sheets
            total_wh_matched = 0
            wh_sheet_names = [s for s in wb.sheetnames if s not in ('ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة')]
            for sname in wh_sheet_names:
                total_wh_matched += max(0, wb[sname].max_row - 1)
                
            curr_not_found = max(0, not_found_ws.max_row - 1) if not_found_ws else 0
            curr_review = max(0, review_ws.max_row - 1) if review_ws.max_row > 1 else 0
            curr_total = total_wh_matched + curr_not_found + curr_review

            # Update the Summary (ملخص) sheet formulas and structure
            if 'ملخص' in wb.sheetnames:
                ws_summary = wb['ملخص']
                num_wh = len(wh_sheet_names)
                ws_summary['B4'] = f"=SUM(B9:B{8 + num_wh})" if num_wh > 0 else "0"
                ws_summary['B5'] = "=MAX(0, COUNTA('لم يُعثر عليه'!A:A)-1)"
                ws_summary['B6'] = "=MAX(0, COUNTA('يحتاج مراجعة'!B:B)-1)"
                
                # Update warehouse detail rows (Row 9+)
                for r in range(9, ws_summary.max_row + 1):
                    wh_title = ws_summary.cell(row=r, column=1).value
                    if wh_title:
                        safe_title = wh_title
                        for ch in r"[]:*?/\\":
                            safe_title = safe_title.replace(ch, "_")
                        safe_title = safe_title[:31]
                        ws_summary.cell(row=r, column=2, value=f"=MAX(0, COUNTA('{safe_title}'!A:A)-1)")
                        ws_summary.cell(row=r, column=3, value=f'=COUNTIF(\'يحتاج مراجعة\'!D:D, "{wh_title}")')

            wb.save(proc['output_file'])
            
            status = 'completed' if curr_review == 0 else 'reviewing'
            self._db.update_process(
                process_id,
                total_shortages=curr_total,
                matched_count=total_wh_matched,
                not_found_count=curr_not_found,
                review_count=curr_review,
                status=status
            )
            
            new_stats = {
                "total": curr_total,
                "total_shortages": curr_total,
                "matched_count": total_wh_matched,
                "not_found_count": curr_not_found,
                "review_count": curr_review,
                "status": status
            }
            
            return {
                "success": True,
                "matched": matched_count,
                "rejected": rejected_count,
                "remaining": curr_review,
                "new_stats": new_stats
            }
            
        except Exception as e:
            return {"success": False, "error": str(e)}

    def rerun_process(self, process_id: str) -> dict:
        with self._worker_lock:
            if self._worker_thread and self._worker_thread.is_alive():
                return {"success": False, "error": "توجد عملية مطابقة قيد التشغيل حالياً، يرجى الانتظار حتى تنتهي"}

            proc = self._db.get_process(process_id)
            if not proc:
                return {"success": False, "error": "العملية غير موجودة"}
            
            proc_dir = Path("data") / "processes" / process_id
            shortage_dir = proc_dir / "input_shortages"
            warehouse_dir = proc_dir / "input_warehouses"
            
            # Collect original input files
            shortage_paths = [str(f) for f in shortage_dir.glob("*") if f.is_file() and not f.name.endswith(".json")]
            warehouse_paths = [str(f) for f in warehouse_dir.glob("*") if f.is_file() and not f.name.endswith(".json")]
            
            if not shortage_paths or not warehouse_paths:
                return {"success": False, "error": "ملفات العملية غير موجودة في المجلد لإعادة تشغيلها"}
                
            self._db.update_process(process_id, status='processing')
            self.add_log(f"--- Starting Re-run for Process [{process_id}] ---")
            
            self._worker_thread = threading.Thread(
                target=self._process_worker,
                args=(process_id, shortage_paths, warehouse_paths, True),
                daemon=True
            )
            self._worker_thread.start()
            return {"success": True, "process_id": process_id}

    def _sync_process_from_excel(self, proc: dict) -> dict:
        if not proc or not proc.get('output_file'):
            return proc
        output_path = proc['output_file']
        if not os.path.exists(output_path):
            return proc
        try:
            wb = openpyxl.load_workbook(output_path, data_only=False)
            snames = wb.sheetnames
            wh_snames = [s for s in snames if s not in ('ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة')]
            
            act_matched = sum(max(0, wb[s].max_row - 1) for s in wh_snames)
            act_not_found = max(0, wb['لم يُعثر عليه'].max_row - 1) if 'لم يُعثر عليه' in snames else 0
            act_review = max(0, wb['يحتاج مراجعة'].max_row - 1) if 'يحتاج مراجعة' in snames else 0
            
            # Auto-repair summary formulas in Excel if outdated
            changed = False
            if 'ملخص' in snames:
                ws_summary = wb['ملخص']
                b6_val = str(ws_summary['B6'].value or '')
                if "'يحتاج مراجعة'!A:A" in b6_val:
                    ws_summary['B6'] = "=MAX(0, COUNTA('يحتاج مراجعة'!B:B)-1)"
                    changed = True
                b5_val = str(ws_summary['B5'].value or '')
                if "'لم يُعثر عليه'!A:A" not in b5_val:
                    ws_summary['B5'] = "=MAX(0, COUNTA('لم يُعثر عليه'!A:A)-1)"
                    changed = True
                if changed:
                    wb.save(output_path)
            
            curr_rev = proc.get('review_count', 0)
            curr_m = proc.get('matched_count', 0)
            curr_nf = proc.get('not_found_count', 0)
            
            if curr_rev != act_review or curr_m != act_matched or curr_nf != act_not_found:
                new_status = 'completed' if act_review == 0 else 'reviewing'
                self._db.update_process(
                    proc['id'],
                    matched_count=act_matched,
                    not_found_count=act_not_found,
                    review_count=act_review,
                    status=new_status
                )
                proc['matched_count'] = act_matched
                proc['not_found_count'] = act_not_found
                proc['review_count'] = act_review
                proc['status'] = new_status
        except Exception as e:
            print(f"[!] Warning: _sync_process_from_excel failed for {proc.get('id')}: {e}")
        return proc

    def get_history(self) -> str:
        procs = self._db.get_all_processes()
        synced = [self._sync_process_from_excel(p) for p in procs]
        return json.dumps(synced)

    def get_process_details(self, process_id: str) -> str:
        proc = self._db.get_process(process_id)
        if proc:
            proc = self._sync_process_from_excel(proc)
            proc['shortage_files'] = json.loads(proc['shortage_files'])
            proc['warehouse_files'] = json.loads(proc['warehouse_files'])
            return json.dumps(proc)
        return "{}"

    def delete_process(self, process_id: str) -> dict:
        try:
            with self._worker_lock:
                self._process_queue = [item for item in self._process_queue if item[0] != process_id]
            self._db.delete_process(process_id)
            return {"success": True}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def open_result_file(self, process_id: str) -> dict:
        try:
            proc = self._db.get_process(process_id)
            if proc and proc.get('output_file') and Path(proc['output_file']).exists():
                os.startfile(proc['output_file'])
                return {"success": True}
            return {"success": False, "error": "ملف النتائج غير موجود"}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def open_result_folder(self, process_id: str) -> dict:
        try:
            proc_dir = self._db.get_process_dir(process_id)
            if proc_dir.exists():
                os.startfile(str(proc_dir))
                return {"success": True}
            return {"success": False, "error": "مجلد العملية غير موجود"}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def open_excel(self, process_id: str) -> dict:
        return self.open_result_file(process_id)

    def open_folder(self, process_id: str) -> dict:
        return self.open_result_folder(process_id)

    def on_closing(self):
        self._server_manager.cleanup()
