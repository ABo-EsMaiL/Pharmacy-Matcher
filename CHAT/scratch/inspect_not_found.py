import openpyxl
from pathlib import Path

file_path = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-001\results_PROC-001.xlsx")
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb["لم يُعثر عليه"]

print(f"Total rows in 'لم يُعثر عليه': {ws.max_row}")
print("\n--- First 25 items in 'لم يُعثر عليه' ---")
for r in range(2, 27):
    print(f"  {ws.cell(r, 1).value}")
