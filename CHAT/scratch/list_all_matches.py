import openpyxl

wb6 = openpyxl.load_workbook(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-006\results_PROC-006.xlsx', data_only=True)

print("="*70)
print("ALL 119 MATCHES IN RESULTS_PROC-006:")
print("="*70)

wh_sheets = [s for s in wb6.sheetnames if s.endswith('.pdf') or s.endswith('.p')]
total = 0
for s in wh_sheets:
    ws = wb6[s]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    print(f"\n--- {s} ({len(rows)} matches) ---")
    for idx, r in enumerate(rows, 1):
        total += 1
        print(f"  {idx:2d}. {r[0]}  ==>  {r[1]}")

print(f"\nTOTAL MATCHES: {total}")
