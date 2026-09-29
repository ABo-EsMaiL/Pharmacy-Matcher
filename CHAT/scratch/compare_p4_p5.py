import openpyxl

def get_wh_matches(file_path):
    wb = openpyxl.load_workbook(file_path)
    wh_sheets = [s for s in wb.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]
    matches = {}
    for ws in wh_sheets:
        sheet = wb[ws]
        matches[ws] = []
        for r in range(2, sheet.max_row + 1):
            req = sheet.cell(r, 1).value
            cand = sheet.cell(r, 2).value
            score = sheet.cell(r, 3).value
            if req and cand:
                matches[ws].append((req.strip(), cand.strip(), str(score or "")))
    return matches

def get_reviews(file_path):
    wb = openpyxl.load_workbook(file_path)
    sheet = wb['يحتاج مراجعة']
    reviews = []
    for r in range(2, sheet.max_row + 1):
        req = sheet.cell(r, 2).value
        cand = sheet.cell(r, 3).value
        wh = sheet.cell(r, 4).value
        reason = sheet.cell(r, 5).value
        if req:
            reviews.append((wh, req, cand, reason))
    return reviews

p4_matches = get_wh_matches(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
p5_matches = get_wh_matches(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-005\results_PROC-005.xlsx")

p4_revs = get_reviews(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
p5_revs = get_reviews(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-005\results_PROC-005.xlsx")

print("=" * 60)
print(f"REVIEWS IN PROC-004 ({len(p4_revs)}):")
for r in p4_revs:
    print(f"  - [{r[0]}] '{r[1]}' vs '{r[2]}' -> {r[3]}")

print("\n" + "=" * 60)
print(f"REVIEWS IN PROC-005 ({len(p5_revs)}):")
for r in p5_revs:
    print(f"  - [{r[0]}] '{r[1]}' vs '{r[2]}' -> {r[3]}")

print("\n" + "=" * 60)
print("WAREHOUSE MATCH COUNTS (PROC-004 vs PROC-005):")
for ws in p4_matches:
    c4 = len(p4_matches[ws])
    c5 = len(p5_matches.get(ws, []))
    print(f"  - {ws}: PROC-004 = {c4} | PROC-005 = {c5} (Diff: {c5 - c4})")

# Look at missing matches in PROC-005
print("\n" + "=" * 60)
print("MATCHES DROPPED IN PROC-005 (were in PROC-004):")
all_p4_items = {}
for ws, items in p4_matches.items():
    for req, cand, score in items:
        all_p4_items[(ws, req)] = (cand, score)

all_p5_items = {}
for ws, items in p5_matches.items():
    for req, cand, score in items:
        all_p5_items[(ws, req)] = (cand, score)

dropped = []
for k, v in all_p4_items.items():
    if k not in all_p5_items:
        dropped.append((k[0], k[1], v[0]))

print(f"Total dropped matches: {len(dropped)}")
for ws, req, cand in dropped:
    print(f"  - [{ws}] '{req}' <---> '{cand}'")

added = []
for k, v in all_p5_items.items():
    if k not in all_p4_items:
        added.append((k[0], k[1], v[0]))

print(f"\nTotal newly added matches: {len(added)}")
for ws, req, cand in added:
    print(f"  - [{ws}] '{req}' <---> '{cand}'")
