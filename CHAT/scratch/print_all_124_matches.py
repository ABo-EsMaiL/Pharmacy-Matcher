import openpyxl
import re

wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
wh_sheets = [
    'جملة العمروووو.pdf', 'السلام شبين.pdf', 'كيور فارما خاص.pdf',
    'مخزن الحياة فارم اسكندريه-170.p', 'الهضبة الاثنين نقدى0.pdf', 'المكرومي المنصورة.pdf'
]

matches_by_wh = {}
for wname in wh_sheets:
    sheet = wb[wname]
    m_list = []
    for r in range(2, sheet.max_row + 1):
        req = sheet.cell(r, 1).value
        cand = sheet.cell(r, 2).value
        score = sheet.cell(r, 3).value
        m_list.append((req, cand, score))
    matches_by_wh[wname] = m_list

for wname, items in matches_by_wh.items():
    print(f"\n{'='*70}")
    print(f"WAREHOUSE: {wname} ({len(items)} matches)")
    print(f"{'='*70}")
    for idx, (req, cand, score) in enumerate(items, 1):
        print(f"{idx:2d}. REQ: '{req}'")
        print(f"    WHS: '{cand}' (Score: {score})")
