import openpyxl

wb3 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-003/results_PROC-003.xlsx')

def get_sheet_data(wb, sheet_name):
    target = None
    for s in wb.sheetnames:
        if s == sheet_name or s.startswith(sheet_name[:15]):
            target = s
            break
    if not target:
        return []
    ws = wb[target]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return []
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    records = []
    for r in rows[1:]:
        d = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        records.append(d)
    return records

for wh in ['السلام شبين.pdf', 'جملة العمروووو.pdf', 'كيور فارما خاص.pdf']:
    wh_matches = get_sheet_data(wb3, wh)
    print(f"\n--- {wh} ({len(wh_matches)} matches) ---")
    for idx, r in enumerate(wh_matches, 1):
        req = r.get('اسم الصنف المطلوب')
        cand = r.get('اسم الصنف في المخزن')
        print(f"  {idx:2d}. المطلوب: '{req}'  <===>  المخزن: '{cand}'")
