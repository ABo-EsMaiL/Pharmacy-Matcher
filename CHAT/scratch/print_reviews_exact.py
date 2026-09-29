import openpyxl

wb1 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-001/results_PROC-001.xlsx')
wb2 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/results_PROC-002.xlsx')

def get_rows(wb, sheet):
    ws = wb[sheet]
    return list(ws.iter_rows(values_only=True))

r1 = get_rows(wb1, 'يحتاج مراجعة')
r2 = get_rows(wb2, 'يحتاج مراجعة')

print('=== PROC-001 (11 Reviews) ===')
for idx, row in enumerate(r1[1:], 1):
    print(f'{idx}. Shortage: {row[1]} | Cand: {row[2]} | WH: {row[3]} | Reason: {row[4]}')

print('\n=== PROC-002 (6 Reviews) ===')
for idx, row in enumerate(r2[1:], 1):
    print(f'{idx}. Shortage: {row[1]} | Cand: {row[2]} | WH: {row[3]} | Reason: {row[4]}')
