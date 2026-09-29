import openpyxl
from pathlib import Path

p6_path = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-006\results_PROC-006.xlsx")
p5_path = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-005\results_PROC-005.xlsx")
p4_path = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")

def inspect_wb(path, label):
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"=== {label} ({path.name}) ===")
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        print(f"  {sheet}: {ws.max_row - 1} rows")
    return wb

wb6 = inspect_wb(p6_path, "PROC-006")
wb5 = inspect_wb(p5_path, "PROC-005")
wb4 = inspect_wb(p4_path, "PROC-004")
