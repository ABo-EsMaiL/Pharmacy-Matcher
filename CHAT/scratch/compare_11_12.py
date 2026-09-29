import openpyxl
from collections import defaultdict

wb11 = openpyxl.load_workbook('data/processes/PROC-011/results_PROC-011.xlsx')
wb12 = openpyxl.load_workbook('data/processes/PROC-012/results_PROC-012.xlsx')

warehouse_sheets = [s for s in wb11.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]

print("="*60)
print("COMPARING PROC-011 vs PROC-012")
print("="*60)

# Compare Matches Per Warehouse
matches_11 = {}
matches_12 = {}

for s in warehouse_sheets:
    m11 = {}
    ws11 = wb11[s]
    for r in list(ws11.rows)[1:]:
        req = str(r[0].value or '').strip()
        wh = str(r[1].value or '').strip()
        conf = str(r[2].value or '').strip()
        if req:
            m11[req] = (wh, conf)
    matches_11[s] = m11

    m12 = {}
    # handle truncated sheet name if any
    ws12 = wb12[s] if s in wb12.sheetnames else None
    if ws12:
        for r in list(ws12.rows)[1:]:
            req = str(r[0].value or '').strip()
            wh = str(r[1].value or '').strip()
            conf = str(r[2].value or '').strip()
            if req:
                m12[req] = (wh, conf)
    matches_12[s] = m12

    set11 = set(m11.keys())
    set12 = set(m12.keys())
    
    common = set11 & set12
    only_11 = set11 - set12
    only_12 = set12 - set11
    
    print(f"\n--- {s} ---")
    print(f"  PROC-011: {len(set11)} matches | PROC-012: {len(set12)} matches")
    print(f"  Common matches: {len(common)}")
    if only_11:
        print(f"  [Removed in PROC-012 / Only in 11] ({len(only_11)}):")
        for req in only_11:
            print(f"    - '{req}' -> '{m11[req][0]}' ({m11[req][1]})")
    if only_12:
        print(f"  [New in PROC-012 / Only in 12] ({len(only_12)}):")
        for req in only_12:
            print(f"    - '{req}' -> '{m12[req][0]}' ({m12[req][1]})")

# Compare Review Sheets
print("\n" + "="*60)
print("COMPARING 'يحتاج مراجعة'")
print("="*60)

rev11 = {}
for r in list(wb11['يحتاج مراجعة'].rows)[1:]:
    req = str(r[1].value or '').strip()
    cand = str(r[2].value or '').strip()
    wh = str(r[3].value or '').strip()
    reason = str(r[4].value or '').strip()
    if req:
        rev11[(req, wh)] = (cand, reason)

rev12 = {}
for r in list(wb12['يحتاج مراجعة'].rows)[1:]:
    req = str(r[1].value or '').strip()
    cand = str(r[2].value or '').strip()
    wh = str(r[3].value or '').strip()
    reason = str(r[4].value or '').strip()
    if req:
        rev12[(req, wh)] = (cand, reason)

set_rev11 = set(rev11.keys())
set_rev12 = set(rev12.keys())

print(f"PROC-011 Review count: {len(set_rev11)}")
print(f"PROC-012 Review count: {len(set_rev12)}")
print(f"Common in Review: {len(set_rev11 & set_rev12)}")

if set_rev11 - set_rev12:
    print(f"\nItems in Review in PROC-011 but NOT in PROC-012 ({len(set_rev11 - set_rev12)}):")
    for req, wh in set_rev11 - set_rev12:
        print(f"  - '{req}' ({wh}) -> Cand: '{rev11[(req, wh)][0]}'")

if set_rev12 - set_rev11:
    print(f"\nItems in Review in PROC-012 but NOT in PROC-011 ({len(set_rev12 - set_rev11)}):")
    for req, wh in set_rev12 - set_rev11:
        print(f"  - '{req}' ({wh}) -> Cand: '{rev12[(req, wh)][0]}' | Reason: {rev12[(req, wh)][1]}")

# Compare Summary Numbers
print("\n" + "="*60)
print("SUMMARY STATS COMPARISON")
print("="*60)
all_matches_11 = set()
for m in matches_11.values():
    all_matches_11.update(m.keys())

all_matches_12 = set()
for m in matches_12.values():
    all_matches_12.update(m.keys())

print(f"Total Unique Shortages Matched in PROC-011: {len(all_matches_11)}")
print(f"Total Unique Shortages Matched in PROC-012: {len(all_matches_12)}")
print(f"Difference in Unique Matched: {len(all_matches_12 - all_matches_11)} new, {len(all_matches_11 - all_matches_12)} dropped")
print(f"New unique matches in PROC-012: {all_matches_12 - all_matches_11}")
print(f"Dropped unique matches in PROC-012: {all_matches_11 - all_matches_12}")
