import openpyxl

wb6 = openpyxl.load_workbook(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-006\results_PROC-006.xlsx', data_only=True)
wb5 = openpyxl.load_workbook(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-005\results_PROC-005.xlsx', data_only=True)

wh_sheets = [s for s in wb6.sheetnames if s.endswith('.pdf') or s.endswith('.p')]

print("=== ACCURATE COMPARISON OF MATCHES (PROC-006 vs PROC-005) ===")

total_m6 = 0
total_m5 = 0

for s in wh_sheets:
    ws6 = wb6[s]
    ws5 = wb5[s]
    
    m6 = [(row[0], row[1]) for row in ws6.iter_rows(min_row=2, values_only=True) if row[0]]
    m5 = [(row[0], row[1]) for row in ws5.iter_rows(min_row=2, values_only=True) if row[0]]
    
    total_m6 += len(m6)
    total_m5 += len(m5)
    
    set6 = set(m6)
    set5 = set(m5)
    
    added = set6 - set5
    removed = set5 - set6
    
    print(f"\n========================================================")
    print(f"Warehouse: {s} | PROC-006: {len(m6)} matches | PROC-005: {len(m5)} matches")
    print(f"========================================================")
    if added:
        print(f"  [+] NEW MATCHES IN PROC-006 ({len(added)}):")
        for req, cand in sorted(added):
            print(f"      • {req}  ==>  {cand}")
    if removed:
        print(f"  [-] REMOVED MATCHES FROM PROC-005 ({len(removed)}):")
        for req, cand in sorted(removed):
            print(f"      • {req}  ==>  {cand}")
    if not added and not removed:
        print("  [=] Identical matches.")

print(f"\n========================================================")
print(f"TOTAL MATCHES: PROC-006 = {total_m6} | PROC-005 = {total_m5}")
print(f"========================================================")
