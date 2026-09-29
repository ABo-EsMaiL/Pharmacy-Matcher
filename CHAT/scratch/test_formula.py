import openpyxl

wb = openpyxl.Workbook()
ws_summary = wb.active
ws_summary.title = 'ملخص'

ws_wh1 = wb.create_sheet('السلام شبين')
ws_wh1.append(['الصنف المطلوب', 'المرشح', 'الدرجة'])
ws_wh1.append(['بنادول', 'بنادول ازرق', 'تام'])
ws_wh1.append(['سيتال', 'سيتال اقراص', 'تام'])

ws_wh2 = wb.create_sheet('جملة العمرو')
ws_wh2.append(['الصنف المطلوب', 'المرشح', 'الدرجة'])
ws_wh2.append(['كليكسان', 'كليكسان 40', 'تام'])

ws_nf = wb.create_sheet('لم يُعثر عليه')
ws_nf.append(['الصنف', 'الملف'])
ws_nf.append(['دواء غير موجود', 'ملف1'])

ws_rev = wb.create_sheet('يحتاج مراجعة')
ws_rev.append(['القرار', 'الصنف المطلوب', 'المرشح', 'المخزن', 'السبب'])
ws_rev.append(['', 'دوليبران 15', 'دوليبران 60', 'السلام شبين', 'فرق اقراص'])

ws_summary.append(['البند', 'العدد', 'النسبة'])
ws_summary.append(['إجمالي الأصناف المطلوبة', '=B3+B4+B5', ''])
ws_summary.append(['تم مطابقتها', '=SUM(B8:B9)', '=IF(B2>0, B3/B2, 0)'])
ws_summary.append(['لم يُعثر عليها', '=MAX(0, COUNTA(\'لم يُعثر عليه\'!A:A)-1)', '=IF(B2>0, B4/B2, 0)'])
ws_summary.append(['تحتاج مراجعة', '=MAX(0, COUNTA(\'يحتاج مراجعة\'!A:A)-1)', '=IF(B2>0, B5/B2, 0)'])
ws_summary.append(['', '', ''])
ws_summary.append(['─── تفاصيل المخازن ───', 'عدد المطابقات', 'أصناف للمراجعة'])
ws_summary.append(['السلام شبين', '=MAX(0, COUNTA(\'السلام شبين\'!A:A)-1)', '=COUNTIF(\'يحتاج مراجعة\'!D:D, "السلام شبين")'])
ws_summary.append(['جملة العمرو', '=MAX(0, COUNTA(\'جملة العمرو\'!A:A)-1)', '=COUNTIF(\'يحتاج مراجعة\'!D:D, "جملة العمرو")'])

wb.save('test_formula.xlsx')
print('Workbook saved successfully with formulas.')
