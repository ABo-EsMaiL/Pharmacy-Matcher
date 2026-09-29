import openpyxl
import os
import re
from rapidfuzz import fuzz

P14_PATH = r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-014\results_PROC-014.xlsx"

import sys
sys.path.insert(0, r"D:\AI_Engineer\Pharmacy-agy")
from src.fast_match.coordinator import (
    get_strength_pattern, has_co, are_forms_identical, extract_form_group,
    extract_brand_tokens, BRAND_FAMILIES
)

def main():
    wb = openpyxl.load_workbook(P14_PATH, data_only=True)
    wh_sheets = [s for s in wb.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]
    
    all_matches = []
    suspicious_matches = []
    
    for s in wh_sheets:
        ws = wb[s]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) <= 1:
            continue
        headers = [str(h or "") for h in rows[0]]
        for r in rows[1:]:
            if not any(r):
                continue
            sh = str(r[0] or "").strip()
            wh_it = str(r[1] or "").strip()
            method = str(r[2] or "").strip() if len(r) > 2 else ""
            
            # Check strength
            s_req = get_strength_pattern(sh)
            s_wh = get_strength_pattern(wh_it)
            str_conflict = False
            if s_req and s_wh and s_req != s_wh:
                str_conflict = True
                
            # Check combo
            co_conflict = (has_co(sh) != has_co(wh_it))
            
            # Check brand sim
            req_b = extract_brand_tokens(sh)
            wh_b = extract_brand_tokens(wh_it)
            b_sim = max(
                fuzz.ratio(req_b, wh_b),
                fuzz.token_sort_ratio(req_b, wh_b)
            ) if req_b and wh_b else 0
            
            # Check forms
            f_req = extract_form_group(sh)
            f_wh = extract_form_group(wh_it)
            form_diff = not are_forms_identical(f_req, f_wh)
            
            all_matches.append({
                "wh": s,
                "shortage": sh,
                "wh_item": wh_it,
                "method": method,
                "str_conflict": str_conflict,
                "s_req": s_req,
                "s_wh": s_wh,
                "co_conflict": co_conflict,
                "b_sim": b_sim,
                "form_diff": form_diff,
                "f_req": f_req,
                "f_wh": f_wh
            })
            
            if str_conflict or co_conflict or b_sim < 65 or form_diff:
                suspicious_matches.append(all_matches[-1])

    print("=" * 80)
    print(f"AUDIT OF ALL {len(all_matches)} MATCHES IN PROC-014")
    print("=" * 80)
    print(f"Total Matches: {len(all_matches)}")
    print(f"Suspicious Matches detected: {len(suspicious_matches)}\n")
    
    for idx, sm in enumerate(suspicious_matches, 1):
        print(f"{idx:02d}. Warehouse: [{sm['wh']}] (Method: {sm['method']})")
        print(f"    Shortage: '{sm['shortage']}'")
        print(f"    WH Item:  '{sm['wh_item']}'")
        issues = []
        if sm['str_conflict']:
            issues.append(f"Strength Conflict: {sm['s_req']} vs {sm['s_wh']}")
        if sm['co_conflict']:
            issues.append("Co/Plus Conflict")
        if sm['b_sim'] < 65:
            issues.append(f"Low Brand Sim: {sm['b_sim']:.1f}%")
        if sm['form_diff']:
            issues.append(f"Form Diff: {sm['f_req']} vs {sm['f_wh']}")
        print(f"    Issues: {', '.join(issues)}\n")

if __name__ == "__main__":
    main()
