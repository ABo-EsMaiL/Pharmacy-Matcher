import openpyxl
import os

P12_PATH = r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-012\results_PROC-012.xlsx"
P14_PATH = r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-014\results_PROC-014.xlsx"

def load_sheet_rows(wb_path, sheet_name):
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h if h is not None else "") for h in rows[0]]
    data = []
    for r in rows[1:]:
        if any(r):
            row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
            row_dict["_raw_tuple"] = r
            data.append(row_dict)
    return headers, data

def main():
    print("=" * 80)
    print("COMPARATIVE AUDIT: PROC-014 vs PROC-012")
    print("=" * 80)
    
    wb12 = openpyxl.load_workbook(P12_PATH, data_only=True)
    wb14 = openpyxl.load_workbook(P14_PATH, data_only=True)
    
    print("\n[1] Sheet Names Comparison:")
    print("PROC-012 sheets:", wb12.sheetnames)
    print("PROC-014 sheets:", wb14.sheetnames)
    
    # Compare per-warehouse matches
    wh_sheets = [s for s in wb14.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]
    
    all_p12_matches = {} # (shortage, wh_item, wh_name)
    all_p14_matches = {}
    
    for s in wh_sheets:
        r12 = load_sheet_rows(P12_PATH, s)
        r14 = load_sheet_rows(P14_PATH, s)
        print(f"\n--- Warehouse: {s} ---")
        print(f"  PROC-012: {len(r12)} matches")
        print(f"  PROC-014: {len(r14)} matches")
        
        m12 = {r.get('اسم الصنف الناقص', ''): r.get('اسم صنف المخزن', '') for r in r12}
        m14 = {r.get('اسم الصنف الناقص', ''): r.get('اسم صنف المخزن', '') for r in r14}
        
        for k, v in m12.items():
            all_p12_matches[(k, s)] = v
        for k, v in m14.items():
            all_p14_matches[(k, s)] = v
            
        diff_dropped = set(m12.keys()) - set(m14.keys())
        diff_added = set(m14.keys()) - set(m12.keys())
        
        if diff_dropped:
            print(f"  [-] In P12 but DROPPED in P14 ({len(diff_dropped)}):")
            for d in diff_dropped:
                print(f"      * Shortage: '{d}' -> P12 WH item: '{m12[d]}'")
        if diff_added:
            print(f"  [+] In P14 but NOT in P12 ({len(diff_added)}):")
            for a in diff_added:
                print(f"      * Shortage: '{a}' -> P14 WH item: '{m14[a]}'")
                
    # Compare Needs Review
    rev12 = load_sheet_rows(P12_PATH, 'يحتاج مراجعة')
    rev14 = load_sheet_rows(P14_PATH, 'يحتاج مراجعة')
    print(f"\n[2] Needs Review Comparison:")
    print(f"  PROC-012: {len(rev12)} items")
    print(f"  PROC-014: {len(rev14)} items")
    
    print("\n--- PROC-014 Needs Review Items ---")
    for r in rev14:
        print(f"  * Shortage: {r.get('اسم الصنف الناقص')} | WH: {r.get('اسم صنف المخزن')} | المخزن: {r.get('المخزن')} | السبب: {r.get('سبب المراجعة')}")

    print("\n--- PROC-012 Needs Review Items ---")
    for r in rev12:
        print(f"  * Shortage: {r.get('اسم الصنف الناقص')} | WH: {r.get('اسم صنف المخزن')} | المخزن: {r.get('المخزن')} | السبب: {r.get('سبب المراجعة')}")

if __name__ == "__main__":
    main()
