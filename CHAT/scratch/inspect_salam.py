import openpyxl
wb = openpyxl.load_workbook(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-001\results_PROC-001.xlsx')
ws = wb['السلام شبين.pdf']
print(f"=== السلام شبين ({len(list(ws.iter_rows()))-1} matches) ===")
for i, r in enumerate(list(ws.iter_rows(values_only=True))[1:], 1):
    print(f"{i:2d}. Req: '{r[0]}' <---> WH: '{r[1]}' [{r[2]}]")
