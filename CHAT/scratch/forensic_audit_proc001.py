import openpyxl

wb = openpyxl.load_workbook(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-001\results_PROC-001.xlsx')

print("="*75)
print("1. SUMMARY SHEET:")
print("="*75)
ws_summary = wb['ملخص']
for r in ws_summary.iter_rows(values_only=True):
    if any(r):
        print(f"  {r[0]:<35} | {str(r[1]):<15} | {str(r[2])}")

print("\n" + "="*75)
print("2. ALL REVIEWS (يحتاج مراجعة):")
print("="*75)
ws_rev = wb['يحتاج مراجعة']
rev_rows = list(ws_rev.iter_rows(values_only=True))[1:]
print(f"Total reviews in sheet: {len(rev_rows)}")
for i, r in enumerate(rev_rows, 1):
    print(f"{i:2d}. [{r[3]}]")
    print(f"    المطلوب: '{r[1]}'")
    print(f"    المخزن : '{r[2]}'")
    print(f"    السبب  : {r[4]}\n")

print("="*75)
print("3. ALL MATCHES PER WAREHOUSE:")
print("="*75)
for sheet in wb.sheetnames:
    if sheet not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))[1:]
        print(f"\n--- {sheet} ({len(rows)} matches) ---")
        for i, r in enumerate(rows, 1):
            print(f"  {i:2d}. Req: '{r[0]:<30}' <---> WH: '{r[1]:<35}' [{r[2]}]")

print("\n" + "="*75)
print("4. NOT FOUND SHEET COUNT:")
print("="*75)
ws_nf = wb['لم يُعثر عليه']
nf_rows = list(ws_nf.iter_rows(values_only=True))[1:]
print(f"Total items in 'لم يُعثر عليه': {len(nf_rows)}")
