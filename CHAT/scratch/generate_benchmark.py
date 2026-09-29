import sys
sys.path.insert(0, r'd:\AI_Engineer\Pharmacy-agy')
import openpyxl, json
from rapidfuzz import fuzz
from src.match.text_match import normalize_text
from src.extract.pdf_reader import parse_structured_result

# 1. Load Shortages (0913.xlsx)
wb_orig = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/input_shortages/0913.xlsx')
ws_orig = wb_orig['ورقة2']

shortages = []
for r in range(7, ws_orig.max_row + 1):
    client = ws_orig.cell(r, 1).value
    name = ws_orig.cell(r, 2).value
    date = ws_orig.cell(r, 3).value
    qty = ws_orig.cell(r, 4).value
    if name:
        shortages.append({
            "row": r,
            "client": client,
            "name": str(name).strip(),
            "date": date,
            "qty": qty
        })

# 2. Load Warehouses
salam_data = json.load(open('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/input_warehouses/السلام شبين.json', encoding='utf-8'))
salam_items = [i['item_name_raw'] for i in parse_structured_result(salam_data, 'salam')]

amr_data = json.load(open('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/input_warehouses/جملة العمروووو.json', encoding='utf-8'))
amr_items = [i['item_name_raw'] for i in parse_structured_result(amr_data, 'amr')]

print(f"Total Shortages: {len(shortages)}")
print(f"Total Salam: {len(salam_items)}")
print(f"Total Amr: {len(amr_items)}")

# 3. Match AMR items
print("\n" + "="*70)
print("AMR (19 ITEMS) MATCH RESULTS IN 0913.xlsx")
print("="*70)
amr_found = []
for idx, a in enumerate(amr_items, 1):
    best_match = None
    best_score = 0
    norm_a = normalize_text(a)
    for sh in shortages:
        norm_sh = normalize_text(sh["name"])
        sc = fuzz.token_set_ratio(norm_a, norm_sh)
        if sc > best_score:
            best_score = sc
            best_match = sh
    print(f"{idx}. Amr: '{a}'")
    if best_score >= 60:
        print(f"   --> MATCHED Row {best_match['row']}: '{best_match['name']}' (Score: {best_score:.1f})")
        amr_found.append((a, best_match, best_score))
    else:
        print(f"   --> NOT FOUND in 0913.xlsx (Closest was Row {best_match['row']}: '{best_match['name']}' with only {best_score:.1f}%)")

print(f"\nTotal Amr items found with confidence in 0913.xlsx: {len(amr_found)} of 19")

# 4. Match SALAM items
print("\n" + "="*70)
print("SALAM (187 ITEMS) MATCH RESULTS IN 0913.xlsx")
print("="*70)
salam_found = []
for idx, s in enumerate(salam_items, 1):
    best_match = None
    best_score = 0
    norm_s = normalize_text(s)
    for sh in shortages:
        norm_sh = normalize_text(sh["name"])
        sc = fuzz.token_set_ratio(norm_s, norm_sh)
        if sc > best_score:
            best_score = sc
            best_match = sh
    if best_score >= 65:
        salam_found.append((idx, s, best_match, best_score))

print(f"Total Salam items found with ratio >= 65%: {len(salam_found)} of 187")
for idx, s, sh, sc in salam_found:
    print(f"[{idx}] Salam: '{s}' ==> Row {sh['row']}: '{sh['name']}' ({sc:.1f}%)")
