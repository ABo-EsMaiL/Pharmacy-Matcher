import sys
sys.path.insert(0, ".")
import openpyxl

wb = openpyxl.load_workbook('data/processes/PROC-009/results_PROC-009.xlsx', data_only=True)
wh_sheets = [s for s in wb.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]

for sheet in wh_sheets:
    ws = wb[sheet]
    print(f"\n=== {sheet} ({ws.max_row - 1} items) ===")
    for r in ws.iter_rows(min_row=2, values_only=True):
        print(f"[{r[2]}]  {r[0]}  --->  {r[1]}")
