import openpyxl

wb3 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-003/results_PROC-003.xlsx')

warehouses = [
    'السلام شبين.pdf',
    'جملة العمروووو.pdf',
    'كيور فارما خاص.pdf',
    'الهضبة الاثنين نقدى0.pdf',
    'المكرومي المنصورة.pdf',
    'مخزن الحياة فارم اسكندريه-170.p'
]

def get_sheet_data(wb, sheet_name):
    target = None
    for s in wb.sheetnames:
        if s == sheet_name or s.startswith(sheet_name[:15]):
            target = s
            break
    if not target:
        return []
    ws = wb[target]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return []
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    records = []
    for r in rows[1:]:
        d = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        records.append(d)
    return records

all_matches = []
for wh in warehouses:
    for r in get_sheet_data(wb3, wh):
        r['_wh'] = wh
        all_matches.append(r)

print(f"Total Matches to audit: {len(all_matches)}")

# Let's inspect each match:
for idx, r in enumerate(all_matches, 1):
    req = r.get('اسم الصنف المطلوب')
    cand = r.get('اسم الصنف في المخزن')
    wh = r['_wh']
    print(f"{idx:3d}. [{wh[:12]}] '{req}' <===> '{cand}'")
