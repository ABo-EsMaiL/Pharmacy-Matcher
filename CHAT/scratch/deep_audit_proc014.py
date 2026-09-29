import openpyxl
import os
import glob

P12_PATH = r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-012\results_PROC-012.xlsx"
P14_PATH = r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-014\results_PROC-014.xlsx"

def load_sheet(wb_path, sheet_name):
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return [], []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h if h is not None else "").strip() for h in rows[0]]
    data = []
    for r in rows[1:]:
        if any(r):
            row_dict = {headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))}
            data.append(row_dict)
    return headers, data

def main():
    print("=" * 80)
    print("DEEP AUDIT & ANALYSIS OF PROC-014 vs PROC-012")
    print("=" * 80)

    # 1. Inspect Sheet Headers
    h14, d14_sample = load_sheet(P14_PATH, "المكرومي المنصورة.pdf")
    print(f"\n[Headers in Warehouse Sheet]: {h14}")
    
    h14_rev, _ = load_sheet(P14_PATH, "يحتاج مراجعة")
    print(f"[Headers in Review Sheet]: {h14_rev}")

    h14_nf, _ = load_sheet(P14_PATH, "لم يُعثر عليه")
    print(f"[Headers in Not Found Sheet]: {h14_nf}")

    # 2. Compare Per-Warehouse Matches
    wb14 = openpyxl.load_workbook(P14_PATH, data_only=True)
    wh_sheets = [s for s in wb14.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]
    
    total_m12 = 0
    total_m14 = 0
    
    print("\n" + "=" * 80)
    print("PER-WAREHOUSE MATCH COMPARISON")
    print("=" * 80)
    
    dropped_overall = []
    added_overall = []

    for s in wh_sheets:
        _, r12 = load_sheet(P12_PATH, s)
        _, r14 = load_sheet(P14_PATH, s)
        total_m12 += len(r12)
        total_m14 += len(r14)
        
        # Key by shortage name
        # Let's inspect column names for shortage and wh item
        col_shortage = [k for k in (r14[0].keys() if r14 else []) if "ناقص" in k or "shortage" in k.lower() or "صنف" in k][0]
        col_wh = [k for k in (r14[0].keys() if r14 else []) if "مخزن" in k or "warehouse" in k.lower()][0]
        
        m12 = {r.get(col_shortage, ''): r.get(col_wh, '') for r in r12}
        m14 = {r.get(col_shortage, ''): r.get(col_wh, '') for r in r14}
        
        print(f"\nWarehouse [{s}]: P12={len(r12)} matches | P14={len(r14)} matches")
        
        diff_dropped = set(m12.keys()) - set(m14.keys())
        diff_added = set(m14.keys()) - set(m12.keys())
        
        if diff_dropped:
            print(f"  [-] Dropped in P14 ({len(diff_dropped)}):")
            for d in sorted(diff_dropped):
                print(f"      * '{d}' -> P12 had: '{m12[d]}'")
                dropped_overall.append((s, d, m12[d]))
                
        if diff_added:
            print(f"  [+] Added in P14 ({len(diff_added)}):")
            for a in sorted(diff_added):
                print(f"      * '{a}' -> P14 matched: '{m14[a]}'")
                added_overall.append((s, a, m14[a]))

    print(f"\nTOTAL MATCHES: PROC-012={total_m12} | PROC-014={total_m14} (Net difference: {total_m14 - total_m12})")
    
    # 3. Inspect Needs Review
    print("\n" + "=" * 80)
    print("NEEDS REVIEW AUDIT (يحتاج مراجعة)")
    print("=" * 80)
    _, rev14 = load_sheet(P14_PATH, "يحتاج مراجعة")
    _, rev12 = load_sheet(P12_PATH, "يحتاج مراجعة")
    
    print(f"PROC-014 Review Items ({len(rev14)}):")
    for idx, r in enumerate(rev14, 1):
        print(f"  {idx}. Shortage: '{r.get('اسم الصنف المطلوب') or r.get('اسم الصنف الناقص') or list(r.values())[0]}' | "
              f"WH: '{r.get('اسم صنف المخزن') or list(r.values())[1]}' | "
              f"WH File: '{r.get('اسم المخزن') or r.get('المخزن') or list(r.values())[2]}' | "
              f"Reason: '{r.get('سبب المراجعة') or list(r.values())[-1]}'")

    print(f"\nPROC-012 Review Items ({len(rev12)}):")
    for idx, r in enumerate(rev12, 1):
        print(f"  {idx}. Shortage: '{list(r.values())[0]}' | WH: '{list(r.values())[1]}' | WH File: '{list(r.values())[2]}' | Reason: '{list(r.values())[-1]}'")

    # 4. Inspect Not Found (لم يُعثر عليه)
    print("\n" + "=" * 80)
    print("NOT FOUND AUDIT (لم يُعثر عليه)")
    print("=" * 80)
    _, nf14 = load_sheet(P14_PATH, "لم يُعثر عليه")
    _, nf12 = load_sheet(P12_PATH, "لم يُعثر عليه")
    print(f"PROC-014 Not Found count: {len(nf14)}")
    print(f"PROC-012 Not Found count: {len(nf12)}")
    
    # Check if any dropped item from P12 is now in P14's Not Found
    nf14_names = set()
    for r in nf14:
        vals = [v for v in r.values() if v is not None]
        if vals:
            nf14_names.add(str(vals[0]).strip())
            
    print(f"\nChecking where the {len(dropped_overall)} dropped matches from P12 went:")
    for wh, s_name, wh_val in dropped_overall:
        in_nf = s_name in nf14_names
        print(f"  - Dropped: '{s_name}' (WH: {wh} -> '{wh_val}') => Is in Not Found: {in_nf}")

if __name__ == "__main__":
    main()
