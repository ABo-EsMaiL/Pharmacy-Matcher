import openpyxl
from pathlib import Path

file_path = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-001\results_PROC-001.xlsx")
wb = openpyxl.load_workbook(file_path, data_only=False)
print("Sheet names:", wb.sheetnames)
ws = wb["ملخص"]
for r in range(1, 16):
    row_vals = [ws.cell(r, c).value for c in range(1, 6)]
    if any(row_vals):
        print(f"Row {r}:", row_vals)

print("\n--- Row counts in sheets ---")
for s in wb.sheetnames:
    print(f"  {s}: {wb[s].max_row} rows")
