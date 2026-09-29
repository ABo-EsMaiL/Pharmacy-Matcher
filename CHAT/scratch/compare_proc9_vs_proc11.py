import sys
sys.path.insert(0, ".")
import openpyxl
from collections import defaultdict
import re

# Load PROC-009 (or PROC-010)
wb_old = openpyxl.load_workbook('data/processes/PROC-009/results_PROC-009.xlsx')
wb_new = openpyxl.load_workbook('data/processes/PROC-011/results_PROC-011.xlsx')

# Extract all matches in PROC-011
new_matches = {} # req -> list of (wh_item, wh_name)
for s in wb_new.sheetnames:
    if s in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']:
        continue
    ws = wb_new[s]
    for r in list(ws.rows)[1:]:
        req = str(r[0].value or '').strip()
        wh_item = str(r[1].value or '').strip()
        if req:
            new_matches.setdefault(req, []).append((wh_item, s))

# Extract all review in PROC-011
new_reviews = {} # req -> list of (cand, wh_name, reason)
ws_rev_new = wb_new['يحتاج مراجعة']
for r in list(ws_rev_new.rows)[1:]:
    req = str(r[1].value or '').strip()
    cand = str(r[2].value or '').strip()
    wh = str(r[3].value or '').strip()
    reason = str(r[4].value or '').strip()
    if req:
        new_reviews.setdefault(req, []).append((cand, wh, reason))

# Extract all not-found in PROC-011
new_not_found = set()
ws_nf_new = wb_new['لم يُعثر عليه']
for r in list(ws_nf_new.rows)[1:]:
    req = str(r[0].value or '').strip()
    if req:
        new_not_found.add(req)

# Now inspect PROC-009 review sheet
ws_rev_old = wb_old['يحتاج مراجعة']
old_review_rows = list(ws_rev_old.rows)[1:]
print(f"Total rows in PROC-009 review sheet: {len(old_review_rows)}")

destination_counts = defaultdict(int)
dest_items = defaultdict(list)

# Track unique shortage items in old review
old_rev_unique_items = set()

for r in old_review_rows:
    req = str(r[1].value or '').strip()
    cand = str(r[2].value or '').strip()
    wh = str(r[3].value or '').strip()
    reason = str(r[4].value or '').strip()
    
    old_rev_unique_items.add(req)
    
    if req in new_matches:
        destination_counts['matched'] += 1
        dest_items['matched'].append((req, cand, wh, reason, new_matches[req]))
    elif req in new_reviews:
        destination_counts['stayed_review'] += 1
        dest_items['stayed_review'].append((req, cand, wh, reason, new_reviews[req]))
    elif req in new_not_found:
        destination_counts['not_found'] += 1
        dest_items['not_found'].append((req, cand, wh, reason))
    else:
        destination_counts['other'] += 1
        dest_items['other'].append((req, cand, wh, reason))

print("\n=== DESTINATION OF OLD 1176 REVIEW ROWS IN PROC-011 ===")
for k, v in destination_counts.items():
    print(f"  {k}: {v} rows")

print(f"\nUnique shortage items in old review: {len(old_rev_unique_items)}")
unq_matched = sum(1 for req in old_rev_unique_items if req in new_matches)
unq_rev = sum(1 for req in old_rev_unique_items if req in new_reviews)
unq_nf = sum(1 for req in old_rev_unique_items if req in new_not_found)
print(f"  -> Unique Matched in PROC-011: {unq_matched}")
print(f"  -> Unique Stayed in Review in PROC-011: {unq_rev}")
print(f"  -> Unique Moved to Not Found in PROC-011: {unq_nf}")

# Let's categorize WHY items moved to Not Found:
from src.fast_match.coordinator import classify_pair

rejection_reasons = defaultdict(int)
sample_rejections = defaultdict(list)

for req, cand, wh, old_reason in dest_items['not_found']:
    status, reason = classify_pair(req, cand)
    rejection_reasons[reason] += 1
    if len(sample_rejections[reason]) < 3:
        sample_rejections[reason].append((req, cand, old_reason))

print("\n=== WHY DID ITEMS MOVE TO NOT FOUND / REJECTED? ===")
for reason, count in sorted(rejection_reasons.items(), key=lambda x: -x[1])[:15]:
    print(f"\nReason: {reason} ({count} occurrences)")
    for req, cand, old_r in sample_rejections[reason]:
        print(f"   * Req: '{req}'")
        print(f"     Cand: '{cand}'")
        print(f"     Old reason: '{old_r}'")
