import openpyxl
from rapidfuzz import fuzz

wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
wh_sheets = [
    'جملة العمروووو.pdf', 'السلام شبين.pdf', 'كيور فارما خاص.pdf',
    'مخزن الحياة فارم اسكندريه-170.p', 'الهضبة الاثنين نقدى0.pdf', 'المكرومي المنصورة.pdf'
]

all_matches = []
for wname in wh_sheets:
    sheet = wb[wname]
    for r in range(2, sheet.max_row + 1):
        req = sheet.cell(r, 1).value
        cand = sheet.cell(r, 2).value
        score = sheet.cell(r, 3).value
        all_matches.append((wname, req, cand, score))

print(f"Loaded {len(all_matches)} matches.")

# Let's inspect potential mismatches
for idx, (wh, req, cand, score) in enumerate(all_matches, 1):
    # compute some metrics
    print(f"[{idx:3d}] ({wh[:15]}) '{req}' <---> '{cand}'")
