import openpyxl
import pandas as pd
from pathlib import Path

p7_path = Path("data/processes/PROC-007/results_PROC-007.xlsx")
p8_path = Path("data/processes/PROC-008/results_PROC-008.xlsx")

wb7 = openpyxl.load_workbook(p7_path, data_only=True)
wb8 = openpyxl.load_workbook(p8_path, data_only=True)

print("=== SUMMARY OF SHEETS ===")
for sheet in wb7.sheetnames:
    s7 = wb7[sheet]
    s8 = wb8[sheet] if sheet in wb8.sheetnames else None
    s8_rows = s8.max_row if s8 else "MISSING"
    print(f"{sheet}: P7={s7.max_row} rows | P8={s8_rows} rows")

print("\n=== DETAILED WAREHOUSE SHEET DIFFS ===")
warehouse_sheets = [s for s in wb7.sheetnames if s not in ["يحتاج مراجعة", "لم يتم العثور عليها"]]

for sname in warehouse_sheets:
    df7 = pd.read_excel(p7_path, sheet_name=sname)
    df8 = pd.read_excel(p8_path, sheet_name=sname)
    
    col_req = "اسم الصنف في النواقص" if "اسم الصنف في النواقص" in df7.columns else df7.columns[0]
    col_wh = "اسم الصنف في المخزن" if "اسم الصنف في المخزن" in df7.columns else df7.columns[1]
    
    items7 = set(df7[col_req].dropna().astype(str).str.strip())
    items8 = set(df8[col_req].dropna().astype(str).str.strip())
    
    only_in_7 = items7 - items8
    only_in_8 = items8 - items7
    
    print(f"\n--- {sname} ---")
    print(f"Total P7: {len(items7)} | Total P8: {len(items8)} | Common: {len(items7 & items8)}")
    if only_in_7:
        print("  Items in P7 but NOT in P8:")
        for it in only_in_7:
            row = df7[df7[col_req].astype(str).str.strip() == it].iloc[0]
            print(f"    - Req: '{it}' -> Wh: '{row[col_wh]}'")
    if only_in_8:
        print("  Items in P8 but NOT in P7:")
        for it in only_in_8:
            row = df8[df8[col_req].astype(str).str.strip() == it].iloc[0]
            print(f"    - Req: '{it}' -> Wh: '{row[col_wh]}'")

print("\n=== REVIEW SHEET DIFFS ===")
df7_rev = pd.read_excel(p7_path, sheet_name="يحتاج مراجعة")
df8_rev = pd.read_excel(p8_path, sheet_name="يحتاج مراجعة")
col_req_r = "اسم الصنف في النواقص" if "اسم الصنف في النواقص" in df7_rev.columns else df7_rev.columns[0]
rev7 = set(df7_rev[col_req_r].dropna().astype(str).str.strip())
rev8 = set(df8_rev[col_req_r].dropna().astype(str).str.strip())

print(f"Total Review P7: {len(rev7)} | Total Review P8: {len(rev8)} | Common: {len(rev7 & rev8)}")
only_in_rev7 = rev7 - rev8
only_in_rev8 = rev8 - rev7
if only_in_rev7:
    print("  In Review P7 but not P8:")
    for it in only_in_rev7:
        row = df7_rev[df7_rev[col_req_r].astype(str).str.strip() == it].iloc[0]
        print(f"    - '{it}' (Warehouse: {row.get('المخزن', '')})")
if only_in_rev8:
    print("  In Review P8 but not P7:")
    for it in only_in_rev8:
        row = df8_rev[df8_rev[col_req_r].astype(str).str.strip() == it].iloc[0]
        print(f"    - '{it}' (Warehouse: {row.get('المخزن', '')})")
