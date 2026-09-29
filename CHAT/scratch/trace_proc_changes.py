import openpyxl

wb1 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-001/results_PROC-001.xlsx')
wb2 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/results_PROC-002.xlsx')

def get_sheet_data(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return []
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    records = []
    for r in rows[1:]:
        d = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        records.append(d)
    return records

print("==================================================")
print("1. SUMMARY SHEET COMPARISON (الملخص)")
print("==================================================")
s1 = get_sheet_data(wb1, 'ملخص')
s2 = get_sheet_data(wb2, 'ملخص')
print("--- PROC-001 ---")
for r in s1:
    print(r)
print("--- PROC-002 ---")
for r in s2:
    print(r)

print("\n==================================================")
print("2. REVIEW ITEMS COMPARISON (يحتاج مراجعة)")
print("==================================================")
r1 = get_sheet_data(wb1, 'يحتاج مراجعة')
r2 = get_sheet_data(wb2, 'يحتاج مراجعة')

print(f"PROC-001 Review items count: {len(r1)}")
print(f"PROC-002 Review items count: {len(r2)}")

print("\n--- PROC-001 ALL 11 REVIEWS ---")
for idx, item in enumerate(r1, 1):
    req = item.get('اسم الصنف المطلوب')
    cand = item.get('اسم المرشح في المخزن')
    wh = item.get('المخزن')
    reason = item.get('السبب')
    print(f"{idx}. [{wh}] '{req}' <--> '{cand}' | السبب: {reason}")

print("\n--- PROC-002 ALL 6 REVIEWS ---")
for idx, item in enumerate(r2, 1):
    req = item.get('اسم الصنف المطلوب')
    cand = item.get('اسم المرشح في المخزن')
    wh = item.get('المخزن')
    reason = item.get('السبب')
    print(f"{idx}. [{wh}] '{req}' <--> '{cand}' | السبب: {reason}")

# Trace the 5 reviews that were in PROC-001 but disappeared in PROC-002:
r2_reqs = {(x.get('اسم الصنف المطلوب'), x.get('المخزن')) for x in r2}
missing_reviews = [x for x in r1 if (x.get('اسم الصنف المطلوب'), x.get('المخزن')) not in r2_reqs]

print(f"\n--- THE 5 REVIEWS THAT CHANGED (In PROC-001 but NOT in PROC-002) ---")
for idx, item in enumerate(missing_reviews, 1):
    req = item.get('اسم الصنف المطلوب')
    cand = item.get('اسم المرشح في المخزن')
    wh = item.get('المخزن')
    reason = item.get('السبب')
    print(f"{idx}. [{wh}] '{req}' <--> '{cand}' | Old Reason: {reason}")

print("\n==================================================")
print("3. WAREHOUSE MATCHES COMPARISON")
print("==================================================")
wh_names = [s for s in wb1.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]

all_m1 = []
all_m2 = []

for wh in wh_names:
    wh_s2 = wh if wh in wb2.sheetnames else [x for x in wb2.sheetnames if x[:15] == wh[:15]][0]
    m1_list = get_sheet_data(wb1, wh)
    m2_list = get_sheet_data(wb2, wh_s2)
    print(f"Warehouse '{wh}': PROC-001 = {len(m1_list)} matches | PROC-002 = {len(m2_list)} matches")
    for r in m1_list:
        r['_wh'] = wh
        all_m1.append(r)
    for r in m2_list:
        r['_wh'] = wh
        all_m2.append(r)

m1_map = {(r.get('اسم الصنف المطلوب'), r['_wh']): r.get('اسم الصنف في المخزن') for r in all_m1}
m2_map = {(r.get('اسم الصنف المطلوب'), r['_wh']): r.get('اسم الصنف في المخزن') for r in all_m2}

dropped_matches = set(m1_map.keys()) - set(m2_map.keys())
new_matches = set(m2_map.keys()) - set(m1_map.keys())
common_matches = set(m1_map.keys()) & set(m2_map.keys())

print(f"\nCommon matches in both: {len(common_matches)}")
print(f"Matches in PROC-001 only (Dropped in PROC-002): {len(dropped_matches)}")
print(f"Matches in PROC-002 only (Newly Added in PROC-002): {len(new_matches)}")

print("\n--- ALL DROPPED MATCHES (Why did they leave?) ---")
for idx, (req, wh) in enumerate(sorted(dropped_matches), 1):
    old_cand = m1_map[(req, wh)]
    # Where did this shortage go in PROC-002? Is it in review, another warehouse, or not found?
    in_r2 = [x for x in r2 if x.get('اسم الصنف المطلوب') == req]
    in_m2_other = [w for (rq, w) in m2_map if rq == req]
    dest = "NOT FOUND"
    if in_r2:
        dest = f"Moved to REVIEW in {in_r2[0].get('المخزن')} (Cand: {in_r2[0].get('اسم المرشح في المخزن')})"
    elif in_m2_other:
        dest = f"Matched in other WH: {in_m2_other}"
    print(f"{idx}. [{wh}] Shortage: '{req}'")
    print(f"    Was matched with: '{old_cand}'")
    print(f"    Current status in PROC-002: {dest}")

print("\n--- ALL NEWLY ADDED MATCHES IN PROC-002 ---")
for idx, (req, wh) in enumerate(sorted(new_matches), 1):
    cand = m2_map[(req, wh)]
    # Where was it in PROC-001?
    in_r1 = [x for x in r1 if x.get('اسم الصنف المطلوب') == req]
    in_m1_other = [w for (rq, w) in m1_map if rq == req]
    origin = "WAS NOT FOUND"
    if in_r1:
        origin = f"Was in REVIEW in {in_r1[0].get('المخزن')} (Cand: {in_r1[0].get('اسم المرشح في المخزن')})"
    elif in_m1_other:
        origin = f"Was matched in other WH: {in_m1_other}"
    print(f"{idx}. [{wh}] Shortage: '{req}'")
    print(f"    Now matched with: '{cand}'")
    print(f"    Previous status in PROC-001: {origin}")
