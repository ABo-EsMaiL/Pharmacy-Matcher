import openpyxl

wb = openpyxl.load_workbook('data/processes/PROC-003/results_PROC-003.xlsx')
ws = wb['السلام شبين.pdf']

print(f"Total rows in السلام شبين: {ws.max_row}")
suspicious = []
for r in range(2, ws.max_row + 1):
    req = str(ws.cell(row=r, column=1).value or '')
    wh = str(ws.cell(row=r, column=2).value or '')
    conf = str(ws.cell(row=r, column=3).value or '')
    
    # Check if conf is llm or if text differs significantly
    if conf == 'llm':
        print(f"[LLM MATCH] Row {r}: Req='{req}' <--> Wh='{wh}'")
    
    # Simple check on words
    req_words = set(req.replace('س ج', '').split())
    wh_words = set(wh.split())
    common = req_words.intersection(wh_words)
    if len(common) == 0:
        suspicious.append((r, req, wh, conf))

if suspicious:
    print(f"\nFound {len(suspicious)} suspicious with 0 common words:")
    for item in suspicious:
        print(item)
else:
    print("\nNo completely disjoint matches found in السلام شبين.")
