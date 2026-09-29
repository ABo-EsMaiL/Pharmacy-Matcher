import openpyxl

wb = openpyxl.load_workbook('data/processes/PROC-003/results_PROC-003.xlsx', data_only=False)
print("Sheet names:", wb.sheetnames)

ws_sum = wb['ملخص']
print("\n--- Summary Sheet (Formulas) ---")
for r in range(1, ws_sum.max_row + 1):
    vals = [ws_sum.cell(row=r, column=c).value for c in range(1, 4)]
    print(f"Row {r:2d}: {vals}")

ws_rev = wb['يحتاج مراجعة']
print(f"\n--- Review Sheet ({ws_rev.max_row} rows) ---")
for r in range(1, ws_rev.max_row + 1):
    vals = [ws_rev.cell(row=r, column=c).value for c in range(1, 6)]
    print(f"Row {r:2d}: {vals}")
