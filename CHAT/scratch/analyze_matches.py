import sys
sys.path.insert(0, r'd:\AI_Engineer\Pharmacy-agy')
import openpyxl, json
from rapidfuzz import fuzz
from src.match.text_match import normalize_text
from src.extract.pdf_reader import parse_structured_result

wb = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/input_shortages/0913.xlsx')
ws = wb['ورقة2']
shortages = []
for r in range(7, ws.max_row + 1):
    c_name = ws.cell(r, 2).value
    if c_name:
        shortages.append((r, str(c_name).strip()))

salam_data = json.load(open('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/input_warehouses/السلام شبين.json', encoding='utf-8'))
salam_items = [i['item_name_raw'] for i in parse_structured_result(salam_data, 'salam')]

amr_data = json.load(open('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/input_warehouses/جملة العمروووو.json', encoding='utf-8'))
amr_items = [i['item_name_raw'] for i in parse_structured_result(amr_data, 'amr')]

print(f"Total shortages in 0913.xlsx: {len(shortages)}")
print(f"Total Salam items: {len(salam_items)}")
print(f"Total Amr items: {len(amr_items)}")

print("\n" + "="*50)
print("--- AMR ITEMS MATCHING ---")
print("="*50)
amr_matches = []
for a in amr_items:
    candidates = []
    norm_a = normalize_text(a)
    for r, s in shortages:
        norm_s = normalize_text(s)
        sc1 = fuzz.token_set_ratio(norm_a, norm_s)
        sc2 = fuzz.token_sort_ratio(norm_a, norm_s)
        sc = max(sc1, sc2)
        if sc >= 50:
            candidates.append((sc, r, s))
    candidates.sort(key=lambda x: x[0], reverse=True)
    if candidates:
        best = candidates[0]
        amr_matches.append((best[0], a, best[1], best[2]))
        print(f"[Score: {best[0]:.1f}] Amr: '{a}'\n  --> Shortage (Row {best[1]}): '{best[2]}'")
    else:
        print(f"[No match >= 50%] Amr: '{a}'")

print("\n" + "="*50)
print("--- SALAM ITEMS MATCHING ---")
print("="*50)
salam_matches = []
for s_item in salam_items:
    candidates = []
    norm_item = normalize_text(s_item)
    for r, s in shortages:
        norm_s = normalize_text(s)
        sc1 = fuzz.token_set_ratio(norm_item, norm_s)
        sc2 = fuzz.token_sort_ratio(norm_item, norm_s)
        sc = max(sc1, sc2)
        if sc >= 50:
            candidates.append((sc, r, s))
    candidates.sort(key=lambda x: x[0], reverse=True)
    if candidates:
        best = candidates[0]
        salam_matches.append((best[0], s_item, best[1], best[2]))

print(f"Salam items with match score >= 50%: {len(salam_matches)}")
salam_matches.sort(key=lambda x: x[0], reverse=True)
for sc, item, r, s in salam_matches[:30]:
    print(f"[Score: {sc:.1f}] Salam: '{item}'\n  --> Shortage (Row {r}): '{s}'")
