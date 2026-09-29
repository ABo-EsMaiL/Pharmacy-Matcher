import openpyxl

file_path = r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx"
wb = openpyxl.load_workbook(file_path)

print("Sheet names:", wb.sheetnames)
for name in wb.sheetnames:
    sheet = wb[name]
    print(f"Sheet '{name}': {sheet.max_row} rows, {sheet.max_column} cols")
    # print headers
    headers = [cell.value for cell in sheet[1]]
    print("  Headers:", headers[:8])
