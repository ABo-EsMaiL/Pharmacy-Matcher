import sys
sys.path.insert(0, 'd:/AI_Engineer/Pharmacy-agy')
import openpyxl
from rapidfuzz import fuzz
from src.config import load_config
from src.extract.pdf_reader import extract_items_from_pdf
from src.match.text_match import normalize_text
from src.fast_match.search import extract_brand_tokens
from pathlib import Path

config = load_config()

wb3 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-003/results_PROC-003.xlsx')

def get_sheet_data(wb, sheet_name):
    target = None
    for s in wb.sheetnames:
        if s == sheet_name or s.startswith(sheet_name[:15]):
            target = s
            break
    if not target:
        return []
    ws = wb[target]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return []
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    records = []
    for r in rows[1:]:
        d = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        records.append(d)
    return records

nf_rows = get_sheet_data(wb3, 'لم يُعثر عليه')
print(f"Total Not Found in PROC-003: {len(nf_rows)}")

wh_files = list(Path('data/processes/PROC-003/input_warehouses').glob('*.pdf'))
all_wh_catalog = []
for p in wh_files:
    items = extract_items_from_pdf(p, config.unstructured_api_key)
    all_wh_catalog.extend(items)

print(f"Total loaded warehouse items: {len(all_wh_catalog)}")

# Find any real brand match that was left in Not Found
missed = []
for row in nf_rows:
    req_name = row.get('اسم الصنف') or list(row.values())[0]
    req_norm = normalize_text(req_name)
    req_brand = extract_brand_tokens(req_name)
    if not req_brand or len(req_brand) < 3:
        continue
    
    for wh_item in all_wh_catalog:
        wh_name = wh_item.get('item_name_raw', '')
        wh_brand = extract_brand_tokens(wh_name)
        if not wh_brand or len(wh_brand) < 3:
            continue
            
        ratio = fuzz.ratio(req_brand, wh_brand)
        if ratio >= 88:
            missed.append({
                'shortage': req_name,
                'wh_item': wh_name,
                'wh_file': wh_item.get('source_file', ''),
                'sim': ratio
            })

seen = set()
unique_missed = []
for m in missed:
    k = (m['shortage'], m['wh_item'])
    if k not in seen:
        seen.add(k)
        unique_missed.append(m)

print(f"\nPotential Missed Matches in Warehouses: {len(unique_missed)}")
for idx, m in enumerate(unique_missed, 1):
    print(f"{idx}. Shortage: '{m['shortage']}' <--> WH: '{m['wh_item']}' [{m['wh_file']}] (sim: {m['sim']}%)")
