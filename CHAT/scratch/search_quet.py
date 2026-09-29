import openpyxl

wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
for name in wb.sheetnames:
    sheet = wb[name]
    for r in range(1, sheet.max_row + 1):
        for c in range(1, sheet.max_column + 1):
            val = str(sheet.cell(r, c).value or "")
            if "كويتابين" in val or "كوتيابين" in val:
                print(f"Found in sheet '{name}' at R{r}C{c}: {val}")
