import openpyxl

wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-014\results_PROC-014.xlsx", data_only=True)
targets = ['نيوروفيت', 'سكينورين', 'نيوكاربون', 'هيرو', 'تربتيزول', 'بقدونس', 'كروم', 'تورسامو']

for name in wb.sheetnames:
    ws = wb[name]
    for r in ws.iter_rows(values_only=True):
        row_str = ' '.join(str(c or '') for c in r)
        for target in targets:
            if target in row_str:
                print(f"Sheet [{name}]: {r}")
