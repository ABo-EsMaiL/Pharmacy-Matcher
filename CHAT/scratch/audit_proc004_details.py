import openpyxl
import json
from collections import defaultdict

wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")

# 1. Review sheet
rev_sheet = wb['يحتاج مراجعة']
print("=" * 60)
print("1. CASES IN REVIEW SHEET (يحتاج مراجعة):")
print("=" * 60)
for r in range(2, rev_sheet.max_row + 1):
    decision = rev_sheet.cell(r, 1).value
    req = rev_sheet.cell(r, 2).value
    cand = rev_sheet.cell(r, 3).value
    wh = rev_sheet.cell(r, 4).value
    reason = rev_sheet.cell(r, 5).value
    print(f"[{r-1}] Shortage: '{req}'")
    print(f"    Candidate: '{cand}' | Warehouse: '{wh}'")
    print(f"    Reason: {reason}\n")

# 2. Matched items across warehouses
wh_sheets = [
    'جملة العمروووو.pdf', 'السلام شبين.pdf', 'كيور فارما خاص.pdf',
    'مخزن الحياة فارم اسكندريه-170.p', 'الهضبة الاثنين نقدى0.pdf', 'المكرومي المنصورة.pdf'
]

all_matches = []
unique_matched_shortages = set()
wh_counts = {}

for wname in wh_sheets:
    sheet = wb[wname]
    count = 0
    for r in range(2, sheet.max_row + 1):
        req = sheet.cell(r, 1).value
        cand = sheet.cell(r, 2).value
        score = sheet.cell(r, 3).value
        all_matches.append({
            "wh": wname,
            "shortage": req,
            "cand": cand,
            "score": score
        })
        unique_matched_shortages.add(req)
        count += 1
    wh_counts[wname] = count

print("=" * 60)
print(f"2. MATCH STATS:")
print(f"Total warehouse match instances: {len(all_matches)}")
print(f"Unique shortage items matched: {len(unique_matched_shortages)}")
for wname, c in wh_counts.items():
    print(f"  - {wname}: {c} matches")
print("=" * 60)

# 3. Not found sheet
not_found_sheet = wb['لم يُعثر عليه']
not_found_items = []
for r in range(2, not_found_sheet.max_row + 1):
    item = not_found_sheet.cell(r, 1).value
    if item:
        not_found_items.append(item)

print(f"3. NOT FOUND ITEMS: {len(not_found_items)} items")
print(f"Total accounted: {len(unique_matched_shortages)} matched + {rev_sheet.max_row - 1} review + {len(not_found_items)} not found = {len(unique_matched_shortages) + rev_sheet.max_row - 1 + len(not_found_items)}")
