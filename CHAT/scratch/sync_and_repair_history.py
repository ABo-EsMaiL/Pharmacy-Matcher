import os
import openpyxl
import sqlite3
from pathlib import Path

def sync_all():
    db_path = Path(r"D:\AI_Engineer\Pharmacy-agy\data\history.db")
    if not db_path.exists():
        print("Database not found.")
        return

    conn = sqlite3.connect(str(db_path))
    c = conn.cursor()
    rows = c.execute('SELECT id, status, total_shortages, matched_count, not_found_count, review_count, output_file FROM processes').fetchall()

    for row in rows:
        pid, st, tot, m, nf, rev, out_f = row
        print(f"Checking {pid}: status={st}, rev_db={rev}, file={out_f}")
        if out_f and os.path.exists(out_f):
            try:
                wb = openpyxl.load_workbook(out_f)
                snames = wb.sheetnames
                wh_snames = [s for s in snames if s not in ('ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة')]
                act_m = sum(max(0, wb[s].max_row - 1) for s in wh_snames)
                act_nf = max(0, wb['لم يُعثر عليه'].max_row - 1) if 'لم يُعثر عليه' in snames else 0
                act_rev = max(0, wb['يحتاج مراجعة'].max_row - 1) if 'يحتاج مراجعة' in snames else 0
                print(f"   -> Actual rows in Excel sheets: matched={act_m}, not_found={act_nf}, review={act_rev}")

                # Check and fix summary sheet formulas
                if 'ملخص' in snames:
                    ws = wb['ملخص']
                    b6_val = str(ws['B6'].value or '')
                    print(f"   -> Cell B6 in summary: {b6_val}")
                    changed = False
                    if "'يحتاج مراجعة'!A:A" in b6_val:
                        ws['B6'] = "=MAX(0, COUNTA('يحتاج مراجعة'!B:B)-1)"
                        changed = True
                    b5_val = str(ws['B5'].value or '')
                    if "'لم يُعثر عليه'!A:A" not in b5_val:
                        ws['B5'] = "=MAX(0, COUNTA('لم يُعثر عليه'!A:A)-1)"
                        changed = True
                    if changed:
                        wb.save(out_f)
                        print(f"   -> Successfully fixed formulas and saved {out_f}")

                # Update database
                if rev != act_rev or m != act_m or nf != act_nf:
                    new_st = 'completed' if act_rev == 0 else 'reviewing'
                    c.execute(
                        'UPDATE processes SET matched_count=?, not_found_count=?, review_count=?, status=? WHERE id=?',
                        (act_m, act_nf, act_rev, new_st, pid)
                    )
                    conn.commit()
                    print(f"   -> Updated DB row for {pid}: review was {rev} -> now {act_rev}")
                else:
                    print(f"   -> DB row was already matching actual Excel count ({act_rev}).")
            except Exception as e:
                print(f"   -> Error processing {pid}: {e}")
        else:
            print("   -> Output file does not exist.")

    conn.close()

if __name__ == "__main__":
    sync_all()
