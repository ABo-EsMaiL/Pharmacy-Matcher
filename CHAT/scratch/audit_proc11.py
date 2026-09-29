import openpyxl

wb = openpyxl.load_workbook('D:/AI_Engineer/Pharmacy-agy/data/processes/PROC-011/results_PROC-011.xlsx')

print("Total sheets:", wb.sheetnames)

total_matches = 0
for s in wb.sheetnames:
    if s in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']:
        continue
    rows = list(wb[s].rows)
    count = len(rows) - 1
    total_matches += count
    print(f"\n=== {s} ({count} matches) ===")
    for r in rows[1:]:
        print(f"  Req:  {r[0].value}")
        print(f"  Wh:   {r[1].value}")
        print(f"  Conf: {r[2].value}")
        print("  " + "-"*40)

print(f"\nTotal matched across all warehouses: {total_matches}")

ws_rev = wb['يحتاج مراجعة']
rev_rows = list(ws_rev.rows)
print(f"\n=== يحتاج مراجعة ({len(rev_rows)-1} items) ===")
for r in rev_rows[1:]:
    print(f"  Req:    {r[1].value}")
    print(f"  Cand:   {r[2].value}")
    print(f"  Wh:     {r[3].value}")
    print(f"  Reason: {r[4].value}")
    print("  " + "-"*40)

ws_nf = wb['لم يُعثر عليه']
nf_rows = list(ws_nf.rows)
print(f"\n=== لم يُعثر عليه ({len(nf_rows)-1} items) ===")
