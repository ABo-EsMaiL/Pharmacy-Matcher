import openpyxl

wb1 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-001/results_PROC-001.xlsx')
wb2 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/results_PROC-002.xlsx')

def get_sheet_data(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return []
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    records = []
    for r in rows[1:]:
        d = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        records.append(d)
    return records

wh_names = [s for s in wb1.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]

for wh in wh_names:
    wh_s2 = wh if wh in wb2.sheetnames else [x for x in wb2.sheetnames if x[:15] == wh[:15]][0]
    m1 = get_sheet_data(wb1, wh)
    m2 = get_sheet_data(wb2, wh_s2)
    
    m1_dict = {r['اسم الصنف المطلوب']: r['اسم الصنف في المخزن'] for r in m1}
    m2_dict = {r['اسم الصنف المطلوب']: r['اسم الصنف في المخزن'] for r in m2}
    
    dropped = set(m1_dict.keys()) - set(m2_dict.keys())
    added = set(m2_dict.keys()) - set(m1_dict.keys())
    common = set(m1_dict.keys()) & set(m2_dict.keys())
    
    print(f"\n{'='*70}")
    print(f"WAREHOUSE: {wh}")
    print(f"PROC-001: {len(m1)} | PROC-002: {len(m2)} | Common: {len(common)} | Dropped: {len(dropped)} | Added: {len(added)}")
    print(f"{'='*70}")
    
    if dropped:
        print("  [-] DROPPED from this warehouse:")
        for k in sorted(dropped):
            print(f"      * '{k}' -> was: '{m1_dict[k]}'")
            
    if added:
        print("  [+] ADDED to this warehouse:")
        for k in sorted(added):
            print(f"      * '{k}' -> now: '{m2_dict[k]}'")
