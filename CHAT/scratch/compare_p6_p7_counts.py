import openpyxl
from pathlib import Path

p7_path = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-007\results_PROC-007.xlsx")
p6_path = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-006\results_PROC-006.xlsx")

def inspect_wb(path, label):
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"=== {label} ({path.name}) ===")
    total_m = 0
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = ws.max_row - 1
        print(f"  {sheet}: {rows} rows")
        if sheet.endswith('.pdf') or sheet.endswith('.p'):
            total_m += rows
    print(f"  TOTAL WAREHOUSE MATCHES: {total_m}\n")
    return wb

wb7 = inspect_wb(p7_path, "PROC-007")
wb6 = inspect_wb(p6_path, "PROC-006")
