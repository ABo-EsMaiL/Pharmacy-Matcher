import openpyxl

wb7 = openpyxl.load_workbook(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-007\results_PROC-007.xlsx', data_only=True)
ws = wb7['يحتاج مراجعة']

print("="*70)
print(f"ALL {ws.max_row - 1} ITEMS IN 'يحتاج مراجعة' (PROC-007):")
print("="*70)

for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 1):
    req = row[1]
    cand = row[2]
    wh = row[3]
    reason = row[4] if len(row) > 4 else ""
    print(f"{i:2d}. المطلوب: {req}")
    print(f"    المرشح:  {cand} ({wh})")
    print(f"    السبب:   {reason}")
