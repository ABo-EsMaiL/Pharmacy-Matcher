import openpyxl
from pathlib import Path
import json

PROCESSES = {
    "PROC-012": Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-012\results_PROC-012.xlsx"),
    "PROC-014": Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-014\results_PROC-014.xlsx"),
    "PROC-015": Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-015\results_PROC-015.xlsx"),
    "PROC-016": Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-016\results_PROC-016.xlsx"),
}

def load_process_excel(proc_id, file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    data = {
        "id": proc_id,
        "sheets": wb.sheetnames,
        "matches": [],
        "reviews": [],
        "not_found": [],
        "matches_by_wh": {},
    }
    
    special_sheets = {"ملخص", "يحتاج مراجعة", "لم يُعثر عليه"}
    
    # 1. Per-warehouse sheets (all sheets that are not special_sheets)
    for name in wb.sheetnames:
        if name in special_sheets:
            continue
        ws = wb[name]
        headers = [str(cell.value or '').strip() for cell in ws[1]]
        wh_matches = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_dict = {headers[i] if i < len(headers) else f"col_{i}": (row[i] or '') for i in range(len(row))}
            row_dict["المخزن"] = name
            wh_matches.append(row_dict)
            data["matches"].append(row_dict)
        data["matches_by_wh"][name] = wh_matches
        
    # 2. Review sheet (يحتاج مراجعة)
    if "يحتاج مراجعة" in wb.sheetnames:
        ws = wb["يحتاج مراجعة"]
        headers = [str(cell.value or '').strip() for cell in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_dict = {headers[i] if i < len(headers) else f"col_{i}": (row[i] or '') for i in range(len(row))}
            data["reviews"].append(row_dict)

    # 3. Not found sheet (لم يُعثر عليه)
    if "لم يُعثر عليه" in wb.sheetnames:
        ws = wb["لم يُعثر عليه"]
        headers = [str(cell.value or '').strip() for cell in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_dict = {headers[i] if i < len(headers) else f"col_{i}": (row[i] or '') for i in range(len(row))}
            data["not_found"].append(row_dict)
            
    return data

def main():
    procs = {}
    for pid, pth in PROCESSES.items():
        if pth.exists():
            procs[pid] = load_process_excel(pid, pth)
            print(f"Loaded {pid}: {len(procs[pid]['matches'])} matches across {len(procs[pid]['matches_by_wh'])} warehouses, {len(procs[pid]['reviews'])} reviews, {len(procs[pid]['not_found'])} not found.")
        else:
            print(f"Warning: {pid} path {pth} does not exist!")

    print("\n" + "="*80)
    print("1. SUMMARY METRICS COMPARISON (المقارنة الرقمية الشاملة للمطابقات حسب المخزن)")
    print("="*80)
    wh_names = sorted(list(set().union(*(p["matches_by_wh"].keys() for p in procs.values()))))
    
    # Print header
    header_str = f"{'المخزن':<30} | " + " | ".join(f"{pid:<10}" for pid in procs.keys())
    print(header_str)
    print("-" * len(header_str))
    
    for wh in wh_names:
        counts = [f"{len(procs[pid]['matches_by_wh'].get(wh, [])):<10}" for pid in procs.keys()]
        print(f"{wh:<30} | " + " | ".join(counts))
        
    print("-" * len(header_str))
    totals = [f"{len(procs[pid]['matches']):<10}" for pid in procs.keys()]
    print(f"{'إجمالي المطابقات':<30} | " + " | ".join(totals))
    
    reviews_str = [f"{len(procs[pid]['reviews']):<10}" for pid in procs.keys()]
    print(f"{'أصناف للمراجعة':<30} | " + " | ".join(reviews_str))

    nf_str = [f"{len(procs[pid]['not_found']):<10}" for pid in procs.keys()]
    print(f"{'لم يتم العثور عليها':<30} | " + " | ".join(nf_str))

    # Save to json for detailed auditing scripts
    export_data = {
        pid: {
            "matches": procs[pid]["matches"],
            "reviews": procs[pid]["reviews"],
            "not_found": procs[pid]["not_found"]
        }
        for pid in procs
    }
    with open(r"C:\Users\kakak\.gemini\antigravity\brain\916274de-7f92-49fc-a0a8-d291664b474a\scratch\extracted_4way_data.json", "w", encoding="utf-8") as f:
        json.dump(export_data, f, ensure_ascii=False, indent=2, default=str)
    print("\nFull data successfully exported to extracted_4way_data.json.")

if __name__ == "__main__":
    main()
