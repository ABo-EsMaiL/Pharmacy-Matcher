import sys
sys.path.insert(0, 'd:/AI_Engineer/Pharmacy-agy')
import openpyxl

wb1 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-001/results_PROC-001.xlsx')
wb2 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/results_PROC-002.xlsx')

def get_sheet_data(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return []
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    records = []
    for r in rows[1:]:
        d = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        records.append(d)
    return records

wh_names = [s for s in wb1.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]

all_m1 = []
all_m2 = []
for wh in wh_names:
    wh_s2 = wh if wh in wb2.sheetnames else [x for x in wb2.sheetnames if x[:15] == wh[:15]][0]
    for r in get_sheet_data(wb1, wh):
        r['_wh'] = wh
        all_m1.append(r)
    for r in get_sheet_data(wb2, wh_s2):
        r['_wh'] = wh
        all_m2.append(r)

r1 = get_sheet_data(wb1, 'يحتاج مراجعة')
r2 = get_sheet_data(wb2, 'يحتاج مراجعة')

m1_map = {(r['اسم الصنف المطلوب'], r['_wh']): r['اسم الصنف في المخزن'] for r in all_m1}
m2_map = {(r['اسم الصنف المطلوب'], r['_wh']): r['اسم الصنف في المخزن'] for r in all_m2}

dropped = set(m1_map.keys()) - set(m2_map.keys())
added = set(m2_map.keys()) - set(m1_map.keys())

print(f"Total Matches in PROC-001: {len(all_m1)}")
print(f"Total Matches in PROC-002: {len(all_m2)}")
print(f"Dropped matches count: {len(dropped)}")
print(f"Added matches count: {len(added)}")

print("\n" + "="*80)
print("ANALYSIS OF DROPPED MATCHES (Were in PROC-001, Removed in PROC-002)")
print("="*80)
for idx, (req, wh) in enumerate(sorted(dropped), 1):
    cand = m1_map[(req, wh)]
    # check if moved to review
    rev_match = [x for x in r2 if x['اسم الصنف المطلوب'] == req and x['المخزن'] == wh]
    # check if matched elsewhere in PROC-002
    other_wh = [w for (rq, w) in m2_map if rq == req]
    status = ""
    if rev_match:
        status = f"[MOVED TO REVIEW in {wh}] -> Reason: {rev_match[0]['السبب']}"
    elif other_wh:
        status = f"[STILL MATCHED in other WH: {other_wh}]"
    else:
        status = "[REJECTED / NOT FOUND]"
    print(f"{idx}. Shortage: '{req}'")
    print(f"   Warehouse: '{wh}' | Previous Match: '{cand}'")
    print(f"   Status in PROC-002: {status}")
    print()

print("\n" + "="*80)
print("ANALYSIS OF ADDED MATCHES (Were NOT in PROC-001, Newly Matched in PROC-002)")
print("="*80)
for idx, (req, wh) in enumerate(sorted(added), 1):
    cand = m2_map[(req, wh)]
    prev_rev = [x for x in r1 if x['اسم الصنف المطلوب'] == req and x['المخزن'] == wh]
    prev_other = [w for (rq, w) in m1_map if rq == req]
    status = ""
    if prev_rev:
        status = f"[PROMOTED FROM REVIEW in {wh}]"
    elif prev_other:
        status = f"[WAS ALREADY MATCHED in other WH: {prev_other}]"
    else:
        status = "[NEW DISCOVERY (Was Not Found)]"
    print(f"{idx}. Shortage: '{req}'")
    print(f"   Warehouse: '{wh}' | New Match: '{cand}'")
    print(f"   Previous in PROC-001: {status}")
    print()
