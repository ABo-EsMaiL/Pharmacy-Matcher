import openpyxl
from pathlib import Path
from collections import Counter

excel_path = Path(r"D:\AI_Engineer\Pharmacy-agy-vision\data\processes\PROC-017\results_PROC-017.xlsx")
wb = openpyxl.load_workbook(excel_path, data_only=True)

print("=== SHEETS IN WORKBOOK ===")
for s in wb.sheetnames:
    ws = wb[s]
    print(f"Sheet: '{s}' -> {ws.max_row} rows, {ws.max_column} cols")

# Summary sheet inspection
if "ملخص" in wb.sheetnames:
    print("\n=== SUMMARY SHEET (ملخص) ===")
    ws_sum = wb["ملخص"]
    for r in range(1, min(15, ws_sum.max_row + 1)):
        row_vals = [ws_sum.cell(r, c).value for c in range(1, min(10, ws_sum.max_column + 1))]
        if any(row_vals):
            print(f"Row {r}: {row_vals}")

# Not found sheet inspection
if "لم يُعثر عليه" in wb.sheetnames:
    ws_nf = wb["لم يُعثر عليه"]
    headers = [ws_nf.cell(1, c).value for c in range(1, ws_nf.max_column + 1)]
    print(f"\n=== NOT FOUND HEADERS: {headers} ===")
    nf_rows = []
    for r in range(2, ws_nf.max_row + 1):
        vals = [ws_nf.cell(r, c).value for c in range(1, ws_nf.max_column + 1)]
        if any(vals):
            nf_rows.append(vals)
    print(f"Total Not Found rows: {len(nf_rows)}")
    print("First 10 Not Found items:")
    for r in nf_rows[:10]:
        print("  *", r)

# Review sheet inspection
if "يحتاج مراجعة" in wb.sheetnames:
    ws_rev = wb["يحتاج مراجعة"]
    headers = [ws_rev.cell(1, c).value for c in range(1, ws_rev.max_column + 1)]
    print(f"\n=== REVIEW HEADERS: {headers} ===")
    rev_rows = []
    for r in range(2, ws_rev.max_row + 1):
        vals = [ws_rev.cell(r, c).value for c in range(1, ws_rev.max_column + 1)]
        if any(vals):
            rev_rows.append(vals)
    print(f"Total Review rows: {len(rev_rows)}")
    print("All Review items:")
    for r in rev_rows:
        print("  *", r)

# Warehouse sheets
wh_sheets = [s for s in wb.sheetnames if s not in ("ملخص", "لم يُعثر عليه", "يحتاج مراجعة")]
print(f"\n=== WAREHOUSE SHEETS ({len(wh_sheets)}) ===")
total_matched = 0
for s in wh_sheets:
    ws = wb[s]
    count = max(0, ws.max_row - 1)
    total_matched += count
    print(f"Warehouse '{s}': {count} matches")
print(f"TOTAL MATCHED ACROSS WAREHOUSES: {total_matched}")
