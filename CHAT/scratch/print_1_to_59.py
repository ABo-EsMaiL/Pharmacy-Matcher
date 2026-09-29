import openpyxl

wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
wh_sheets = [
    'جملة العمروووو.pdf', 'السلام شبين.pdf', 'كيور فارما خاص.pdf',
    'مخزن الحياة فارم اسكندريه-170.p', 'الهضبة الاثنين نقدى0.pdf', 'المكرومي المنصورة.pdf'
]

all_matches = []
for wname in wh_sheets:
    sheet = wb[wname]
    for r in range(2, sheet.max_row + 1):
        req = sheet.cell(r, 1).value
        cand = sheet.cell(r, 2).value
        score = sheet.cell(r, 3).value
        all_matches.append((wname, req, cand, score))

for idx in range(0, min(60, len(all_matches))):
    wh, req, cand, score = all_matches[idx]
    print(f"[{idx+1:3d}] ({wh[:15]}) '{req}' <---> '{cand}'")
