import openpyxl

f1 = 'd:/AI_Engineer/Pharmacy-agy/data/processes/PROC-001/results_PROC-001.xlsx'
f2 = 'd:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/results_PROC-002.xlsx'
f3 = 'd:/AI_Engineer/Pharmacy-agy/data/processes/PROC-003/results_PROC-003.xlsx'

wb1 = openpyxl.load_workbook(f1)
wb2 = openpyxl.load_workbook(f2)
wb3 = openpyxl.load_workbook(f3)

print("=== SHEET NAMES IN PROC-003 ===")
print(wb3.sheetnames)

def get_sheet_data(wb, sheet_name):
    # find matching sheet
    target = None
    for s in wb.sheetnames:
        if s == sheet_name or s.startswith(sheet_name[:15]):
            target = s
            break
    if not target:
        return []
    ws = wb[target]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return []
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    records = []
    for r in rows[1:]:
        d = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        records.append(d)
    return records

warehouses = [
    'السلام شبين.pdf',
    'جملة العمروووو.pdf',
    'كيور فارما خاص.pdf',
    'الهضبة الاثنين نقدى0.pdf',
    'المكرومي المنصورة.pdf',
    'مخزن الحياة فارم اسكندريه-170.p'
]

print("\n" + "="*85)
print(f"{'البيان / المخزن':<32} | {'PROC-001':<12} | {'PROC-002':<12} | {'PROC-003':<12}")
print("="*85)

tot1 = tot2 = tot3 = 0
for wh in warehouses:
    c1 = len(get_sheet_data(wb1, wh))
    c2 = len(get_sheet_data(wb2, wh))
    c3 = len(get_sheet_data(wb3, wh))
    tot1 += c1
    tot2 += c2
    tot3 += c3
    print(f"{wh:<32} | {c1:<12} | {c2:<12} | {c3:<12}")

print("-" * 85)
print(f"{'إجمالي المطابقات (Matches)':<32} | {tot1:<12} | {tot2:<12} | {tot3:<12}")

r1 = len(get_sheet_data(wb1, 'يحتاج مراجعة'))
r2 = len(get_sheet_data(wb2, 'يحتاج مراجعة'))
r3 = len(get_sheet_data(wb3, 'يحتاج مراجعة'))
print(f"{'يحتاج مراجعة (Reviews)':<32} | {r1:<12} | {r2:<12} | {r3:<12}")

nf1 = len(get_sheet_data(wb1, 'لم يُعثر عليه'))
nf2 = len(get_sheet_data(wb2, 'لم يُعثر عليه'))
nf3 = len(get_sheet_data(wb3, 'لم يُعثر عليه'))
print(f"{'لم يُعثر عليه (Not Found)':<32} | {nf1:<12} | {nf2:<12} | {nf3:<12}")
print("=" * 85)
