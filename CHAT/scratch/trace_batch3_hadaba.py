import sys
sys.path.insert(0, r"D:\AI_Engineer\Pharmacy-agy")
from src.extract.excel_reader import read_excel, extract_item_names
from src.extract.pdf_reader import parse_structured_result
from src.fast_match.search import FastCandidateFinder
import json
from pathlib import Path

# Load shortages
excel_path = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\input_shortages\0913.xlsx")
shortages_raw = extract_item_names(read_excel(excel_path))
unique_shortages = {}
for item in shortages_raw:
    name = item.get("item_name_raw", "")
    if name and name not in unique_shortages:
        unique_shortages[name] = item
deduped_shortages = list(unique_shortages.values())

cache_file = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\data\.unstructured_cache\a4f524fd75cf223f927a228ad87319f0_extract.json")
data = json.loads(cache_file.read_text(encoding="utf-8"))
items = parse_structured_result(data, "الهضبة الاثنين نقدى0.pdf")
wh_items = [{"id": f"W_{i}", "name": it["item_name_raw"]} for i, it in enumerate(items)]

finder = FastCandidateFinder()
finder.fit(wh_items)

llm_cases = []
for idx, s_item in enumerate(deduped_shortages):
    name = s_item["item_name_raw"]
    cands = finder.search(name, top_k=6)
    if cands:
        llm_cases.append({
            "case_id": f"Q{idx:05}",
            "requested_name": name,
            "candidates": [{"id": c["id"], "name": c["name"]} for c in cands[:3]]
        })

batch_3 = llm_cases[150:225]

# Which items in batch 3 were matched in PROC-004?
import openpyxl
wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
sheet = wb['الهضبة الاثنين نقدى0.pdf']
p4_hadaba_matches = [sheet.cell(r, 1).value for r in range(2, sheet.max_row + 1)]

wb5 = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-005\results_PROC-005.xlsx")
sheet5 = wb5['الهضبة الاثنين نقدى0.pdf']
p5_hadaba_matches = [sheet5.cell(r, 1).value for r in range(2, sheet5.max_row + 1)]

print(f"Total cases in Batch 3: {len(batch_3)}")
print("\nItems in Batch 3 that were MATCHED in PROC-004:")
b3_p4 = []
for c in batch_3:
    req = c['requested_name']
    if req in p4_hadaba_matches:
        in_p5 = "STILL IN P5" if req in p5_hadaba_matches else "DROPPED IN P5"
        print(f"  - [{in_p5}] '{req}' vs '{c['candidates'][0]['name']}'")
        b3_p4.append(req)

print(f"\nTotal matched in Batch 3 in PROC-004: {len(b3_p4)}")
