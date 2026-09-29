import openpyxl

file1 = 'd:/AI_Engineer/Pharmacy-agy/data/processes/PROC-001/results_PROC-001.xlsx'
file2 = 'd:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/results_PROC-002.xlsx'

wb1 = openpyxl.load_workbook(file1)
wb2 = openpyxl.load_workbook(file2)

def extract_rows(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    data = []
    for r in rows[1:]:
        row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        data.append(row_dict)
    return data

print("=== PROC-001 SUMMARY ===")
for r in extract_rows(wb1, 'ملخص'):
    print(r)

print("\n=== PROC-002 SUMMARY ===")
for r in extract_rows(wb2, 'ملخص'):
    print(r)

print("\n=== PROC-001 ALL 11 REVIEWS ===")
for idx, r in enumerate(extract_rows(wb1, 'يحتاج مراجعة'), 1):
    print(f"{idx}. Shortage: {r.get('صنف النواقص') or r.get('اسم الصنف المطلوب')} | Match: {r.get('صنف المخزن المطابق') or r.get('الصنف المطابق')} | Warehouse: {r.get('المخزن')} | Reason: {r.get('سبب المراجعة') or r.get('السبب')}")

print("\n=== PROC-002 ALL 6 REVIEWS ===")
for idx, r in enumerate(extract_rows(wb2, 'يحتاج مراجعة'), 1):
    print(f"{idx}. Shortage: {r.get('صنف النواقص') or r.get('اسم الصنف المطلوب')} | Match: {r.get('صنف المخزن المطابق') or r.get('الصنف المطابق')} | Warehouse: {r.get('المخزن')} | Reason: {r.get('سبب المراجعة') or r.get('السبب')}")

# All 29 dropped matches
wh_sheets = [s for s in wb1.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]
m1_total = []
m2_total = []
for s in wh_sheets:
    m1_rows = extract_rows(wb1, s)
    s2_match = s if s in wb2.sheetnames else [x for x in wb2.sheetnames if x[:15] == s[:15]][0]
    m2_rows = extract_rows(wb2, s2_match)
    for r in m1_rows:
        r['_wh'] = s
        m1_total.append(r)
    for r in m2_rows:
        r['_wh'] = s
        m2_total.append(r)

m1_keys = {}
for r in m1_total:
    shortage = r.get('صنف النواقص') or r.get('اسم الصنف المطلوب') or r.get('اسم الصنف')
    cand = r.get('صنف المخزن المطابق') or r.get('الصنف المطابق')
    m1_keys[(shortage, r['_wh'])] = (cand, r)

m2_keys = {}
for r in m2_total:
    shortage = r.get('صنف النواقص') or r.get('اسم الصنف المطلوب') or r.get('اسم الصنف')
    cand = r.get('صنف المخزن المطابق') or r.get('الصنف المطابق')
    m2_keys[(shortage, r['_wh'])] = (cand, r)

dropped_matches = set(m1_keys.keys()) - set(m2_keys.keys())
new_matches = set(m2_keys.keys()) - set(m1_keys.keys())

print(f"\n=== ALL 29 DROPPED MATCHES (Present in PROC-001, Missing in PROC-002) ===")
for idx, k in enumerate(sorted(dropped_matches), 1):
    cand, r = m1_keys[k]
    print(f"{idx}. Shortage: '{k[0]}' | Old Match: '{cand}' | Warehouse: {k[1]}")

print(f"\n=== ALL 24 NEW MATCHES (Missing in PROC-001, Present in PROC-002) ===")
for idx, k in enumerate(sorted(new_matches), 1):
    cand, r = m2_keys[k]
    print(f"{idx}. Shortage: '{k[0]}' | New Match: '{cand}' | Warehouse: {k[1]}")
