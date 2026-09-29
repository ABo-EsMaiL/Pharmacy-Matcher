import openpyxl

wb = openpyxl.load_workbook('data/processes/PROC-003/results_PROC-003.xlsx')
print('Sheet names:', wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n{'='*20} {sheetname} ({ws.max_row} rows) {'='*20}")
    limit = 20 if sheetname != 'ملخص' else 30
    for r in range(1, min(limit, ws.max_row + 1)):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"Row {r:2d}: {row_vals}")
