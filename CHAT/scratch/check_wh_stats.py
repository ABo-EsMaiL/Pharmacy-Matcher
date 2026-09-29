import openpyxl

wb1 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-001/results_PROC-001.xlsx')
wb2 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/results_PROC-002.xlsx')

print(f"{'المخزن':<30} | {'PROC-001':<10} | {'PROC-002':<10} | {'الفرق':<10}")
print("-" * 65)

for s in wb1.sheetnames:
    if s in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']:
        continue
    s2 = s if s in wb2.sheetnames else [x for x in wb2.sheetnames if x[:15] == s[:15]][0]
    c1 = wb1[s].max_row - 1
    c2 = wb2[s2].max_row - 1
    diff = c2 - c1
    print(f"{s:<30} | {c1:<10} | {c2:<10} | {diff:+d}")

print("-" * 65)
rev1 = wb1['يحتاج مراجعة'].max_row - 1
rev2 = wb2['يحتاج مراجعة'].max_row - 1
print(f"{'يحتاج مراجعة':<30} | {rev1:<10} | {rev2:<10} | {rev2 - rev1:+d}")

nf1 = wb1['لم يُعثر عليه'].max_row - 1
nf2 = wb2['لم يُعثر عليه'].max_row - 1
print(f"{'لم يُعثر عليه':<30} | {nf1:<10} | {nf2:<10} | {nf2 - nf1:+d}")
