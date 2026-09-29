import openpyxl

wb3 = openpyxl.load_workbook('d:/AI_Engineer/Pharmacy-agy/data/processes/PROC-003/results_PROC-003.xlsx')

def get_sheet_data(wb, sheet_name):
    target = None
    for s in wb.sheetnames:
        if s == sheet_name or s.startswith(sheet_name[:15]):
            target = s
            break
    if not target:
        return []
    ws = wb[target]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return []
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    records = []
    for r in rows[1:]:
        d = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        records.append(d)
    return records

warehouses = [
    'السلام شبين.pdf',
    'جملة العمروووو.pdf',
    'كيور فارما خاص.pdf',
    'الهضبة الاثنين نقدى0.pdf',
    'المكرومي المنصورة.pdf',
    'مخزن الحياة فارم اسكندريه-170.p'
]

print("="*90)
print("1. ALL 8 REVIEWS IN PROC-003 (يحتاج مراجعة)")
print("="*90)
r3 = get_sheet_data(wb3, 'يحتاج مراجعة')
for idx, r in enumerate(r3, 1):
    req = r.get('اسم الصنف المطلوب')
    cand = r.get('اسم المرشح في المخزن')
    wh = r.get('المخزن')
    reason = r.get('السبب')
    print(f"{idx}. [{wh}]")
    print(f"   المطلوب: '{req}'")
    print(f"   المرشح:  '{cand}'")
    print(f"   السبب:   {reason}")
    print()

print("="*90)
print("2. ALL 100 MATCHES IN PROC-003 (مفهرسة حسب المخزن)")
print("="*90)

all_matches = []
for wh in warehouses:
    wh_matches = get_sheet_data(wb3, wh)
    print(f"\n--- {wh} ({len(wh_matches)} matches) ---")
    for idx, r in enumerate(wh_matches, 1):
        req = r.get('اسم الصنف المطلوب')
        cand = r.get('اسم الصنف في المخزن')
        conf = r.get('درجة التطابق')
        all_matches.append({'wh': wh, 'req': req, 'cand': cand, 'conf': conf})
        print(f"  {idx:2d}. المطلوب: '{req}'  <===>  المخزن: '{cand}'")

print(f"\nTotal matches audited: {len(all_matches)}")
