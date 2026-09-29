import os
import glob
import openpyxl

def inspect_all():
    procs = sorted(glob.glob(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-*'))
    for proc in procs:
        pname = os.path.basename(proc)
        excel_files = glob.glob(os.path.join(proc, '*.xlsx'))
        shortages = [os.path.basename(s) for s in glob.glob(os.path.join(proc, 'input_shortages', '*'))]
        whs = [os.path.basename(w) for w in glob.glob(os.path.join(proc, 'input_warehouses', '*'))]
        
        counts = {}
        for ef in excel_files:
            try:
                wb = openpyxl.load_workbook(ef, read_only=True)
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    counts[sheet] = ws.max_row
            except Exception as e:
                counts["error"] = str(e)
                
        print(f"{pname}: Shortages={shortages} | WHs count={len(whs)} | Excel={counts}")

if __name__ == "__main__":
    inspect_all()
