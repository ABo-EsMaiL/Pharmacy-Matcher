import openpyxl
from pathlib import Path

file_path = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-001\results_PROC-001.xlsx")
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb["يحتاج مراجعة"]

print(f"Total rows in 'يحتاج مراجعة': {ws.max_row}")
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
print("Headers:", headers)

print("\n--- First 20 items in 'يحتاج مراجعة' ---")
for r in range(2, 22):
    row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    print(f"Row {r}: {row_vals}")

# Group by reason or pattern
from collections import Counter
reasons = [ws.cell(r, 6).value for r in range(2, ws.max_row + 1)]
print("\n--- Common reasons ---")
for reason, count in Counter(reasons).most_common(10):
    print(f"  [{count} times] {reason}")
