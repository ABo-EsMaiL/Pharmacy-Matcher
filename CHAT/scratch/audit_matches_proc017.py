import openpyxl
from pathlib import Path

excel_path = Path(r"D:\AI_Engineer\Pharmacy-agy-vision\data\processes\PROC-017\results_PROC-017.xlsx")
wb = openpyxl.load_workbook(excel_path, data_only=True)

wh_sheets = [s for s in wb.sheetnames if s not in ("ملخص", "لم يُعثر عليه", "يحتاج مراجعة")]

all_matches = []
for s in wh_sheets:
    ws = wb[s]
    for r in range(2, ws.max_row + 1):
        sh_val = str(ws.cell(r, 1).value or "").strip()
        wh_val = str(ws.cell(r, 2).value or "").strip()
        method = str(ws.cell(r, 3).value or "").strip()
        if sh_val and wh_val:
            all_matches.append({
                "shortage": sh_val,
                "warehouse": wh_val,
                "sheet": s,
                "method": method
            })

print(f"Total matches loaded: {len(all_matches)}")

# Print by warehouse
for s in wh_sheets:
    sheet_matches = [m for m in all_matches if m["sheet"] == s]
    print(f"\n==================================================")
    print(f"WAREHOUSE: {s} ({len(sheet_matches)} matches)")
    print(f"==================================================")
    for idx, m in enumerate(sheet_matches, 1):
        print(f"{idx:2d}. Shortage: '{m['shortage']}' <==> Warehouse: '{m['warehouse']}' [{m['method']}]")
