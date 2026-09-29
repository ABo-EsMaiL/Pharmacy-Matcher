import sys
import os
import openpyxl
from pathlib import Path

# Add project root to sys.path
PROJECT_ROOT = Path("d:/AI_Engineer/Pharmacy-agy")
sys.path.insert(0, str(PROJECT_ROOT))

from desktop_app.backend import API

print("[1] Initializing API...")
api = API()

# Check concurrency check
status = api.is_process_running()
print(f"is_process_running: {status}")

print("[2] Triggering rerun_process for PROC-003...")
res = api.rerun_process("PROC-003")
print(f"rerun_process result: {res}")

if not res.get("success"):
    print("Failed to start rerun:", res)
    sys.exit(1)

# Wait for worker thread to finish
print("[3] Waiting for worker thread to finish matching...")
if api._worker_thread:
    api._worker_thread.join()
print("[+] Worker thread finished.")

# Check concurrency check again
status = api.is_process_running()
print(f"is_process_running after finish: {status}")

# Check the generated Excel file
excel_file = Path("data/processes/PROC-003/results_PROC-003.xlsx")
if not excel_file.exists():
    print("Error: Output Excel not found!")
    sys.exit(1)

wb = openpyxl.load_workbook(excel_file)
print(f"\n[4] Inspecting sheets in {excel_file.name}: {wb.sheetnames}")

# Inspect Summary sheet
ws_sum = wb['ملخص']
print("\n--- Summary Sheet Cells & Formulas ---")
for r in range(1, ws_sum.max_row + 1):
    vals = [ws_sum.cell(row=r, column=c).value for c in range(1, 4)]
    print(f"Row {r:2d}: {vals}")

# Inspect a warehouse sheet to check Arabic confidence
for sname in wb.sheetnames:
    if sname not in ('ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة'):
        ws_wh = wb[sname]
        print(f"\n--- First 5 rows of {sname} (Confidence column) ---")
        for r in range(1, min(6, ws_wh.max_row + 1)):
            print(f"Row {r}: {[ws_wh.cell(row=r, column=c).value for c in range(1, 4)]}")

# Inspect Review sheet
ws_rev = wb['يحتاج مراجعة']
print(f"\n--- Review Sheet ({ws_rev.max_row} rows) ---")
for r in range(1, ws_rev.max_row + 1):
    print(f"Row {r}: {[ws_rev.cell(row=r, column=c).value for c in range(1, 6)]}")

# Now test apply_reviews:
# Set decision on first review item to '✅ مطابق'
if ws_rev.max_row > 1:
    print("\n[5] Simulating user review decision: Setting Row 2 to '✅ مطابق'...")
    ws_rev.cell(row=2, column=1, value='✅ مطابق')
    wb.save(excel_file)
    
    print("[6] Calling api.apply_reviews('PROC-003')...")
    apply_res = api.apply_reviews('PROC-003')
    print(f"apply_reviews result: {apply_res}")
    
    # Reload and inspect summary sheet again
    wb2 = openpyxl.load_workbook(excel_file)
    ws_sum2 = wb2['ملخص']
    print("\n--- Summary Sheet after apply_reviews ---")
    for r in range(1, ws_sum2.max_row + 1):
        vals = [ws_sum2.cell(row=r, column=c).value for c in range(1, 4)]
        print(f"Row {r:2d}: {vals}")

print("\n[+] Verification script completed successfully.")
