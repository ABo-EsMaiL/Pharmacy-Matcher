import openpyxl

wb = openpyxl.load_workbook(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-001\results_PROC-001.xlsx')

for sheet in wb.sheetnames:
    if sheet not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))[1:]
        print(f"\n==================== {sheet} ({len(rows)} matches) ====================")
        for i, r in enumerate(rows, 1):
            print(f"{i:2d}. Req: '{r[0]}' <---> WH: '{r[1]}' [{r[2]}]")
