import openpyxl

file1 = 'd:/AI_Engineer/Pharmacy-agy/data/processes/PROC-001/results_PROC-001.xlsx'
file2 = 'd:/AI_Engineer/Pharmacy-agy/data/processes/PROC-002/results_PROC-002.xlsx'

wb1 = openpyxl.load_workbook(file1)
wb2 = openpyxl.load_workbook(file2)

print("=== SHEET NAMES ===")
print("PROC-001 sheets:", wb1.sheetnames)
print("PROC-002 sheets:", wb2.sheetnames)

def extract_rows(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    data = []
    for r in rows[1:]:
        row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        data.append(row_dict)
    return data

# Let's inspect Summary sheet
print("\n=== SUMMARY SHEET COMPARISON ===")
s1 = extract_rows(wb1, 'الملخص والإحصائيات')
s2 = extract_rows(wb2, 'الملخص والإحصائيات')
print("PROC-001 Summary:")
for row in s1:
    print(" ", row)
print("PROC-002 Summary:")
for row in s2:
    print(" ", row)

# Compare Matches
m1 = extract_rows(wb1, 'مطابقة صحيحة')
m2 = extract_rows(wb2, 'مطابقة صحيحة')
print(f"\nPROC-001 Matched rows: {len(m1)}")
print(f"PROC-002 Matched rows: {len(m2)}")

# Compare Reviews
r1 = extract_rows(wb1, 'يحتاج مراجعة')
r2 = extract_rows(wb2, 'يحتاج مراجعة')
print(f"\nPROC-001 Review rows: {len(r1)}")
print(f"PROC-002 Review rows: {len(r2)}")

# Compare Not Found
nf1 = extract_rows(wb1, 'غير متوفر')
nf2 = extract_rows(wb2, 'غير متوفر')
print(f"\nPROC-001 Not Found rows: {len(nf1)}")
print(f"PROC-002 Not Found rows: {len(nf2)}")
