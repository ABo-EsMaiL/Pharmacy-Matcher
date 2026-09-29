import openpyxl
import re

excel_path = r"D:\AI_Engineer\Pharmacy-agy-vision\data\processes\PROC-017\results_PROC-017.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
wh_sheets = [s for s in wb.sheetnames if s not in ("ملخص", "لم يُعثر عليه", "يحتاج مراجعة")]

suspicious = []
for s in wh_sheets:
    ws = wb[s]
    for r in range(2, ws.max_row + 1):
        sh = str(ws.cell(r, 1).value or "").strip()
        wh = str(ws.cell(r, 2).value or "").strip()
        method = str(ws.cell(r, 3).value or "").strip()
        
        # Check numbers in shortage vs warehouse
        nums_sh = re.findall(r'\d+(?:\.\d+)?', sh.replace('١','1').replace('٢','2').replace('٣','3').replace('٤','4').replace('٥','5').replace('٦','6').replace('٧','7').replace('٨','8').replace('٩','9').replace('٠','0'))
        nums_wh = re.findall(r'\d+(?:\.\d+)?', wh.replace('١','1').replace('٢','2').replace('٣','3').replace('٤','4').replace('٥','5').replace('٦','6').replace('٧','7').replace('٨','8').replace('٩','9').replace('٠','0'))
        
        # Check specific critical words: 1 vs 2 (e.g. hero baby 1 vs 2), dolphin 50 vs 25
        suspicious_reason = []
        if ("هيرو" in sh or "بيبي" in sh or "لبن" in sh) and ("1" in nums_sh or "2" in nums_sh) and ("1" in nums_wh or "2" in nums_wh):
            if set(nums_sh) & {"1", "2"} != set(nums_wh) & {"1", "2"}:
                suspicious_reason.append("اختلاف مرحلة لبن أطفال (1 مقابل 2)")
                
        if "دولفين" in sh and ("50" in nums_sh or "12.5" in nums_sh) and "25" in nums_wh:
            suspicious_reason.append("اختلاف تركيز دولفين (50 أو 12.5 مقابل 25)")
            
        if "500" in nums_sh and "1000" in nums_wh or "1" in nums_sh and "500" in nums_wh:
            if "سيتال" in sh:
                suspicious_reason.append("اختلاف تركيز سيتال")

        if suspicious_reason:
            suspicious.append({
                "warehouse": s,
                "shortage": sh,
                "warehouse_item": wh,
                "method": method,
                "reason": " + ".join(suspicious_reason)
            })

print(f"Found {len(suspicious)} suspicious matches:")
for sp in suspicious:
    print(f"[{sp['warehouse']}] Shortage: '{sp['shortage']}' vs WH: '{sp['warehouse_item']}'")
    print(f"   Method: {sp['method']} | Reason: {sp['reason']}")
    print("-" * 50)
