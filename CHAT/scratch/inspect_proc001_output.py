import openpyxl

wb = openpyxl.load_workbook(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-001\results_PROC-001.xlsx')
print('Sheets:', wb.sheetnames)

ws_summary = wb['ملخص']
print('\n=== ملخص ===')
for r in ws_summary.iter_rows(values_only=True):
    if any(r):
        print(r)

for name in wb.sheetnames:
    if name != 'ملخص':
        ws = wb[name]
        rows = [r for r in ws.iter_rows(values_only=True) if any(r)]
        print(f"\n=== Sheet: '{name}' | Total non-empty rows: {len(rows)} ===")
        if rows:
            print(' Header:', rows[0])
            for r in rows[1:4]:
                print('  Row:', r)
