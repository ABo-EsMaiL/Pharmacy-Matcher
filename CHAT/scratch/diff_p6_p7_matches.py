import openpyxl

wb7 = openpyxl.load_workbook(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-007\results_PROC-007.xlsx', data_only=True)
wb6 = openpyxl.load_workbook(r'D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-006\results_PROC-006.xlsx', data_only=True)

wh_sheets = [s for s in wb7.sheetnames if s.endswith('.pdf') or s.endswith('.p')]

print("="*70)
print("CHANGES IN WAREHOUSE MATCHES (PROC-007 vs PROC-006):")
print("="*70)

for s in wh_sheets:
    ws7 = wb7[s]
    ws6 = wb6[s]
    
    m7 = [(row[0], row[1]) for row in ws7.iter_rows(min_row=2, values_only=True) if row[0]]
    m6 = [(row[0], row[1]) for row in ws6.iter_rows(min_row=2, values_only=True) if row[0]]
    
    set7 = set(m7)
    set6 = set(m6)
    
    added = set7 - set6
    removed = set6 - set7
    
    print(f"\n--- {s} (PROC-007: {len(m7)} | PROC-006: {len(m6)}) ---")
    if added:
        print(f"  [+] NEW MATCHES IN PROC-007 ({len(added)}):")
        for req, cand in sorted(added):
            print(f"      • {req}  ==>  {cand}")
    if removed:
        print(f"  [-] REMOVED MATCHES IN PROC-007 ({len(removed)}):")
        for req, cand in sorted(removed):
            print(f"      • {req}  ==>  {cand}")
    if not added and not removed:
        print("  [=] Identical matches.")
