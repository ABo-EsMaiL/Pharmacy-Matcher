import sys
sys.path.insert(0, r"D:\AI_Engineer\Pharmacy-agy")
import openpyxl
from pathlib import Path
import json
import re
from rapidfuzz import fuzz
from src.extract.pdf_reader import parse_structured_result
from src.fast_match.search import extract_brand_tokens, extract_form_group
from src.match.text_match import normalize_text

wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
sheet = wb['لم يُعثر عليه']
not_found_list = [sheet.cell(r, 1).value.strip() for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value]

wh_folder = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\input_warehouses")
cache_dir = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\data\.unstructured_cache")

import hashlib
def get_hash(fp):
    h = hashlib.md5()
    with open(fp, 'rb') as f:
        while b := f.read(65536):
            h.update(b)
    return h.hexdigest()

all_wh_items = []
for pdf in wh_folder.glob("*.pdf"):
    fhash = get_hash(pdf)
    cfile = cache_dir / f"{fhash}_extract.json"
    if cfile.exists():
        with open(cfile, "r", encoding="utf-8") as f:
            data = json.load(f)
        items = parse_structured_result(data, pdf.name)
        for it in items:
            raw_name = it.get("item_name_raw", "").strip()
            if raw_name:
                all_wh_items.append({
                    "wh_name": pdf.name,
                    "item_raw": raw_name,
                    "brand": extract_brand_tokens(raw_name),
                    "norm": normalize_text(raw_name),
                    "form": extract_form_group(raw_name)
                })

candidates = []

for idx, req in enumerate(not_found_list, 1):
    req_b = extract_brand_tokens(req)
    req_f = extract_form_group(req)
    if not req_b or len(req_b) < 3:
        continue
        
    for wh in all_wh_items:
        wh_raw = wh["item_raw"]
        wh_b = wh["brand"]
        wh_f = wh["form"]
        if not wh_b or len(wh_b) < 3:
            continue
            
        # Avoid prefix match if prefix is very short
        if min(len(req_b), len(wh_b)) <= 3 and (req_b.startswith(wh_b) or wh_b.startswith(req_b)):
            sim = fuzz.ratio(req_b, wh_b)
        else:
            sim = fuzz.ratio(req_b, wh_b)
            if len(req_b.split()) > 1 or len(wh_b.split()) > 1:
                sim = max(sim, fuzz.token_sort_ratio(req_b, wh_b))
            if len(req_b) >= 4 and len(wh_b) >= 4 and (req_b.startswith(wh_b) or wh_b.startswith(req_b)):
                sim = max(sim, 85)
                
        if sim >= 75:
            # Check form
            if req_f and wh_f and req_f != wh_f:
                continue
            candidates.append({
                "idx": idx,
                "shortage": req,
                "cand": wh_raw,
                "wh": wh["wh_name"],
                "sim": sim,
                "req_b": req_b,
                "wh_b": wh_b,
                "req_f": req_f,
                "wh_f": wh_f
            })

# Sort by similarity descending
candidates.sort(key=lambda x: x["sim"], reverse=True)

# Group by shortage
grouped = {}
for c in candidates:
    s = c["shortage"]
    if s not in grouped:
        grouped[s] = []
    grouped[s].append(c)

print(f"Total shortages with viable matches: {len(grouped)}")
for s, c_list in list(grouped.items())[:30]:
    print(f"\n[Shortage] '{s}'")
    for c in c_list[:2]:
        print(f"   -> ({c['wh']}): '{c['cand']}' (Sim: {c['sim']:.1f}, Brand: '{c['req_b']}' vs '{c['wh_b']}')")
