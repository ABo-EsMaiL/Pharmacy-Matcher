import openpyxl

wb = openpyxl.load_workbook('data/processes/PROC-009/results_PROC-009.xlsx', data_only=True)
print('Sheets in PROC-009:', wb.sheetnames)

ws_sum = wb['ملخص']
for r in ws_sum.iter_rows(values_only=True):
    print(r)

ws_rev = wb['يحتاج مراجعة']
print('\nTotal review rows in sheet:', ws_rev.max_row - 1)
print('\nSample 30 review rows:')
for i, r in enumerate(ws_rev.iter_rows(min_row=2, max_row=31, values_only=True)):
    print(f'{i+1}: {r[1]} -> {r[2]} | {r[3]} | {r[4]}')

ws_nf = wb['لم يُعثر عليه']
print('\nTotal not found rows in sheet:', ws_nf.max_row - 1)
