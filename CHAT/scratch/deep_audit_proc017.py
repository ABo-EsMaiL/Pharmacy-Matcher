import json
import openpyxl
from pathlib import Path
from rapidfuzz import fuzz

proc_dir = Path(r"D:\AI_Engineer\Pharmacy-agy-vision\data\processes\PROC-017")
excel_path = proc_dir / "results_PROC-017.xlsx"

wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Load all warehouse items from cache
wh_items = {}
cache_dir = proc_dir / "cache" / "vision"

for wh_folder in cache_dir.iterdir():
    if not wh_folder.is_dir():
        continue
    res_dir = wh_folder / "results"
    if not res_dir.exists():
        continue
    folder_name = wh_folder.name
    wh_items[folder_name] = []
    for p_json in sorted(res_dir.glob("page_*.json")):
        with open(p_json, "r", encoding="utf-8") as fp:
            p_data = json.load(fp)
            items = p_data.get("items", [])
            for it in items:
                wh_items[folder_name].append(it)

total_wh_items = sum(len(v) for v in wh_items.values())
print(f"Loaded {total_wh_items} warehouse items across {len(wh_items)} warehouse folders.")

# 2. Load Matched items
wh_sheets = [s for s in wb.sheetnames if s not in ("ملخص", "لم يُعثر عليه", "يحتاج مراجعة")]
matched_pairs = []
for s in wh_sheets:
    ws = wb[s]
    for r in range(2, ws.max_row + 1):
        sh_val = ws.cell(r, 1).value
        wh_val = ws.cell(r, 2).value
        method = ws.cell(r, 3).value
        if sh_val and wh_val:
            matched_pairs.append({
                "shortage": sh_val,
                "warehouse_item": wh_val,
                "warehouse": s,
                "method": method
            })

print(f"Total matched pairs: {len(matched_pairs)}")

# 3. Load Review items
rev_items = []
if "يحتاج مراجعة" in wb.sheetnames:
    ws = wb["يحتاج مراجعة"]
    for r in range(2, ws.max_row + 1):
        sh_val = ws.cell(r, 2).value
        wh_val = ws.cell(r, 3).value
        wh_name = ws.cell(r, 4).value
        reason = ws.cell(r, 5).value
        if sh_val:
            rev_items.append({
                "shortage": sh_val,
                "warehouse_item": wh_val,
                "warehouse": wh_name,
                "reason": reason
            })

print(f"Total review items: {len(rev_items)}")

# 4. Load Not Found items
nf_items = []
if "لم يُعثر عليه" in wb.sheetnames:
    ws = wb["لم يُعثر عليه"]
    for r in range(2, ws.max_row + 1):
        sh_val = ws.cell(r, 1).value
        src = ws.cell(r, 2).value
        if sh_val:
            nf_items.append(sh_val)

print(f"Total Not Found items: {len(nf_items)}")

# 5. Build flat list of all warehouse extracted items
all_wh_raw = []
for folder, it_list in wh_items.items():
    for it in it_list:
        raw_name = it.get("item_name_raw", "")
        trade = it.get("trade_name", "")
        strength = it.get("strength", "")
        form = it.get("form", "")
        all_wh_raw.append({
            "folder": folder,
            "raw": raw_name,
            "trade": trade,
            "strength": strength,
            "form": form,
            "full_search": f"{trade} {strength} {form}".strip() or raw_name
        })

# 6. Deep audit of Not Found items against Warehouse inventory
print("\n--- AUDITING NOT FOUND ITEMS (Checking for high similarity candidates) ---")

borderline_cases = []
absent_count = 0
partial_match_count = 0

for nf in nf_items:
    best_score = 0
    best_cand = None
    for cand in all_wh_raw:
        s1 = fuzz.token_sort_ratio(nf, cand["raw"])
        s2 = fuzz.token_set_ratio(nf, cand["raw"])
        score = max(s1, s2)
        if score > best_score:
            best_score = score
            best_cand = cand
    
    if best_score >= 80:
        borderline_cases.append({
            "not_found": nf,
            "best_cand": best_cand["raw"],
            "score": best_score,
            "warehouse": best_cand["folder"],
            "trade": best_cand["trade"],
            "strength": best_cand["strength"],
            "form": best_cand["form"]
        })
    elif best_score >= 60:
        partial_match_count += 1
    else:
        absent_count += 1

print(f"High similarity candidates (score >= 80) in Not Found: {len(borderline_cases)}")
print(f"Partial similarity (60 <= score < 80): {partial_match_count}")
print(f"Completely absent from inventory (score < 60): {absent_count}")

print("\n--- DETAILED HIGH SIMILARITY NOT FOUND CASES ---")
for bc in borderline_cases:
    print(f"Shortage: '{bc['not_found']}'")
    print(f"  Candidate: '{bc['best_cand']}' ({bc['warehouse']}) [Score: {bc['score']}]")
    print(f"  Parsed: Trade='{bc['trade']}', Strength='{bc['strength']}', Form='{bc['form']}'")
    print("-" * 50)
