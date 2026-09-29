import sys
sys.path.insert(0, r"D:\AI_Engineer\Pharmacy-agy")
import openpyxl
from pathlib import Path
import json
import re
from rapidfuzz import fuzz
from src.extract.pdf_reader import parse_structured_result
from src.fast_match.search import extract_brand_tokens, extract_form_group
from src.match.text_match import normalize_text

# 1. Load 485 Not Found items
wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
sheet = wb['لم يُعثر عليه']
not_found_list = [sheet.cell(r, 1).value.strip() for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value]

# 2. Load 2,717 warehouse items
wh_folder = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\input_warehouses")
cache_dir = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\data\.unstructured_cache")

import hashlib
def get_hash(fp):
    h = hashlib.md5()
    with open(fp, 'rb') as f:
        while b := f.read(65536):
            h.update(b)
    return h.hexdigest()

all_wh_items = []
for pdf in wh_folder.glob("*.pdf"):
    fhash = get_hash(pdf)
    cfile = cache_dir / f"{fhash}_extract.json"
    if cfile.exists():
        with open(cfile, "r", encoding="utf-8") as f:
            data = json.load(f)
        items = parse_structured_result(data, pdf.name)
        for it in items:
            raw_name = it.get("item_name_raw", "").strip()
            if raw_name:
                all_wh_items.append({
                    "wh_name": pdf.name,
                    "item_raw": raw_name,
                    "brand": extract_brand_tokens(raw_name),
                    "norm": normalize_text(raw_name),
                    "form": extract_form_group(raw_name)
                })

def extract_strength(text):
    m = re.search(r'(?:^|[^\d/])(\d+(?:\.\d+)?\/\d+(?:\.\d+)?)(?:[^\d/]|$)', text)
    if m:
        return m.group(1)
    m = re.search(r'(\d+(?:\.\d+)?)\s*(?:مجم|مج|ملجم|جم|gm|mg)', text)
    if m:
        return m.group(1)
    return None

true_missed = []
brand_matches_different_form_or_str = []

for idx, req in enumerate(not_found_list, 1):
    req_b = extract_brand_tokens(req)
    req_f = extract_form_group(req)
    req_s = extract_strength(req)
    
    if not req_b:
        continue
        
    for wh in all_wh_items:
        wh_raw = wh["item_raw"]
        wh_b = wh["brand"]
        wh_f = wh["form"]
        wh_s = extract_strength(wh_raw)
        
        if not wh_b:
            continue
            
        # check brand similarity
        b_sim = fuzz.ratio(req_b, wh_b)
        if len(req_b.split()) > 1 or len(wh_b.split()) > 1:
            b_sim = max(b_sim, fuzz.token_sort_ratio(req_b, wh_b))
        if req_b.startswith(wh_b) or wh_b.startswith(req_b):
            b_sim = max(b_sim, 85)
            
        if b_sim >= 80:
            # Check form compatibility
            form_ok = (req_f == wh_f) or (not req_f) or (not wh_f)
            
            # Check strength compatibility
            str_ok = True
            if req_s and wh_s:
                if req_s != wh_s:
                    str_ok = False
                    
            # Check combination
            req_co = bool(re.search(r'\b(كو|بلس|بلاس|كومب|co|plus|comp)\b', req.lower()))
            wh_co = bool(re.search(r'\b(كو|بلس|بلاس|كومب|co|plus|comp)\b', wh_raw.lower()))
            co_ok = (req_co == wh_co)
            
            # Check new
            req_new = bool(re.search(r'\b(نيو|new)\b', req.lower()))
            wh_new = bool(re.search(r'\b(نيو|new)\b', wh_raw.lower()))
            new_ok = (req_new == wh_new)
            
            if form_ok and str_ok and co_ok and new_ok:
                true_missed.append({
                    "idx": idx,
                    "shortage": req,
                    "req_b": req_b,
                    "req_f": req_f,
                    "req_s": req_s,
                    "wh_name": wh["wh_name"],
                    "cand": wh_raw,
                    "wh_b": wh_b,
                    "wh_f": wh_f,
                    "wh_s": wh_s,
                    "b_sim": b_sim
                })
            else:
                reason = []
                if not form_ok: reason.append(f"Form ({req_f} vs {wh_f})")
                if not str_ok: reason.append(f"Strength ({req_s} vs {wh_s})")
                if not co_ok: reason.append(f"Combination ({req_co} vs {wh_co})")
                if not new_ok: reason.append(f"New prefix ({req_new} vs {wh_new})")
                brand_matches_different_form_or_str.append({
                    "shortage": req,
                    "cand": wh_raw,
                    "wh": wh["wh_name"],
                    "reasons": ", ".join(reason)
                })

# Deduplicate true_missed
dedup_missed = {}
for m in true_missed:
    key = (m["shortage"], m["cand"], m["wh_name"])
    if key not in dedup_missed:
        dedup_missed[key] = m

print(f"\n==================================================")
print(f"POTENTIAL TRUE MISSED MATCHES: {len(dedup_missed)}")
print(f"==================================================")
for k, m in dedup_missed.items():
    print(f"REQ: '{m['shortage']}'")
    print(f"  -> WH ({m['wh_name']}): '{m['cand']}' (Brand Sim: {m['b_sim']})\n")

print(f"\nTotal legitimate rejections with similar brand: {len(brand_matches_different_form_or_str)}")
