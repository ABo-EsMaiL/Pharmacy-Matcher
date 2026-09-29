import openpyxl
from collections import Counter
from pathlib import Path

wb = openpyxl.load_workbook(r"d:\AI_Engineer\Pharmacy-agy\data\processes\PROC-001\input_shortages\0913.xlsx")
ws = wb["ورقة2"]
drugs = [ws.cell(r, 2).value for r in range(7, ws.max_row+1) if ws.cell(r, 2).value]
print(f"Total raw rows: {len(drugs)}")
unique = list(set(drugs))
print(f"Unique distinct drug names: {len(unique)}")
counts = Counter(drugs)
print("\nTop 10 repeated drugs:")
for d, c in counts.most_common(10):
    print(f"  [{c} times] {d}")
