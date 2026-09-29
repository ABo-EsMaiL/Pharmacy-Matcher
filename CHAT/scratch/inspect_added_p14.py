import openpyxl

P12_PATH = r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-012\results_PROC-012.xlsx"
P14_PATH = r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-014\results_PROC-014.xlsx"

def load_wh_matches(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    wh_sheets = [s for s in wb.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]
    matches = {} # (wh_sheet, shortage) -> wh_item
    for s in wh_sheets:
        ws = wb[s]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) > 1:
            for r in rows[1:]:
                if any(r):
                    shortage = str(r[0] or "").strip()
                    wh_item = str(r[1] or "").strip()
                    matches[(s, shortage)] = wh_item
    return matches

def main():
    m12 = load_wh_matches(P12_PATH)
    m14 = load_wh_matches(P14_PATH)
    
    added_keys = set(m14.keys()) - set(m12.keys())
    dropped_keys = set(m12.keys()) - set(m14.keys())
    
    print(f"Total Added in P14: {len(added_keys)}")
    print(f"Total Dropped in P14: {len(dropped_keys)}")
    
    print("\n" + "=" * 80)
    print("ALL ADDED ITEMS IN PROC-014 (New matches found):")
    print("=" * 80)
    for wh, sh in sorted(added_keys):
        print(f"Warehouse: [{wh}]")
        print(f"   Shortage: '{sh}'")
        print(f"   P14 Match: '{m14[(wh, sh)]}'\n")

if __name__ == "__main__":
    main()
