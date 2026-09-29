import openpyxl

wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
wh_sheets = [
    'جملة العمروووو.pdf', 'السلام شبين.pdf', 'كيور فارما خاص.pdf',
    'مخزن الحياة فارم اسكندريه-170.p'
]

for wname in wh_sheets:
    sheet = wb[wname]
    print(f"\n{'='*70}")
    print(f"WAREHOUSE: {wname} ({sheet.max_row - 1} matches)")
    print(f"{'='*70}")
    for r in range(2, sheet.max_row + 1):
        req = sheet.cell(r, 1).value
        cand = sheet.cell(r, 2).value
        score = sheet.cell(r, 3).value
        print(f"{r-1:2d}. REQ: '{req}'")
        print(f"    WHS: '{cand}' (Score: {score})")
