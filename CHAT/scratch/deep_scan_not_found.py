import sys
sys.path.insert(0, r"D:\AI_Engineer\Pharmacy-agy")
import openpyxl
from pathlib import Path
import json
import re
from rapidfuzz import fuzz
from src.extract.pdf_reader import parse_structured_result
from src.fast_match.search import extract_brand_tokens, extract_form_group
from src.match.text_match import normalize_text

# 1. Load 485 Not Found items
wb = openpyxl.load_workbook(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")
sheet = wb['لم يُعثر عليه']
not_found_list = []
for r in range(2, sheet.max_row + 1):
    val = sheet.cell(r, 1).value
    if val:
        not_found_list.append(val.strip())

print(f"Total not-found items to scan: {len(not_found_list)}")

# 2. Load 2,717 warehouse items
wh_folder = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\input_warehouses")
cache_dir = Path(r"D:\AI_Engineer\Pharmacy-agy\data\processes\data\.unstructured_cache")

import hashlib
def get_hash(fp):
    h = hashlib.md5()
    with open(fp, 'rb') as f:
        while b := f.read(65536):
            h.update(b)
    return h.hexdigest()

all_wh_items = []
for pdf in wh_folder.glob("*.pdf"):
    fhash = get_hash(pdf)
    cfile = cache_dir / f"{fhash}_extract.json"
    if cfile.exists():
        with open(cfile, "r", encoding="utf-8") as f:
            data = json.load(f)
        items = parse_structured_result(data, pdf.name)
        for it in items:
            raw_name = it.get("item_name_raw", "").strip()
            if raw_name:
                all_wh_items.append({
                    "wh_name": pdf.name,
                    "item_raw": raw_name,
                    "brand": extract_brand_tokens(raw_name),
                    "norm": normalize_text(raw_name),
                    "form": extract_form_group(raw_name)
                })

print(f"Total warehouse items indexed: {len(all_wh_items)}")

# 3. Deep Scan
missed_candidates = []
close_candidates = []

for idx, req in enumerate(not_found_list, 1):
    req_brand = extract_brand_tokens(req)
    req_norm = normalize_text(req)
    req_form = extract_form_group(req)
    
    # search across all warehouse items
    matches_for_req = []
    for wh in all_wh_items:
        wh_raw = wh["item_raw"]
        wh_brand = wh["brand"]
        
        # 1. Exact brand substring
        if req_brand and wh_brand:
            if req_brand == wh_brand:
                score = 100
            elif len(req_brand) >= 4 and (req_brand in wh_brand or wh_brand in req_brand):
                score = 90
            else:
                score = max(
                    fuzz.ratio(req_brand, wh_brand),
                    fuzz.token_sort_ratio(req_brand, wh_brand)
                )
        else:
            score = fuzz.token_set_ratio(req_norm, wh["norm"])
            
        if score >= 70:
            matches_for_req.append((wh["wh_name"], wh_raw, score, wh["form"]))
            
    if matches_for_req:
        matches_for_req.sort(key=lambda x: x[2], reverse=True)
        # Check top matches
        best_wh, best_item, best_score, best_form = matches_for_req[0]
        close_candidates.append({
            "idx": idx,
            "shortage": req,
            "req_brand": req_brand,
            "req_form": req_form,
            "best_wh": best_wh,
            "best_cand": best_item,
            "best_score": best_score,
            "best_form": best_form,
            "all_top": matches_for_req[:3]
        })

print(f"\nScan Complete!")
print(f"Items with potential warehouse mentions (score >= 70): {len(close_candidates)} / {len(not_found_list)}")
print(f"Items with ZERO mentions across all 6 PDFs: {len(not_found_list) - len(close_candidates)} / {len(not_found_list)}")
