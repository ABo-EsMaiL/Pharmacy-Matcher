import openpyxl

def get_matches(file_path):
    wb = openpyxl.load_workbook(file_path)
    wh_sheets = [s for s in wb.sheetnames if s not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']]
    matches = {}
    for ws in wh_sheets:
        sheet = wb[ws]
        for r in range(2, sheet.max_row + 1):
            req = sheet.cell(r, 1).value
            cand = sheet.cell(r, 2).value
            if req and cand:
                matches[(ws, req.strip())] = cand.strip()
    return matches

p3_matches = get_matches(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-003\results_PROC-003.xlsx")
p4_matches = get_matches(r"D:\AI_Engineer\Pharmacy-agy\data\processes\PROC-004\results_PROC-004.xlsx")

print(f"Total matches in PROC-003: {len(p3_matches)}")
print(f"Total matches in PROC-004: {len(p4_matches)}")

new_in_p4 = {k: v for k, v in p4_matches.items() if k not in p3_matches}
removed_in_p4 = {k: v for k, v in p3_matches.items() if k not in p4_matches}
changed_in_p4 = {k: (p3_matches[k], p4_matches[k]) for k in p4_matches if k in p3_matches and p4_matches[k] != p3_matches[k]}

print(f"\nNew in PROC-004: {len(new_in_p4)} items")
for idx, ((wh, req), cand) in enumerate(new_in_p4.items(), 1):
    print(f"[{idx:2d}] Warehouse: '{wh}'")
    print(f"     Shortage: '{req}'")
    print(f"     Candidate: '{cand}'\n")

print(f"\nRemoved in PROC-004 (were in PROC-003): {len(removed_in_p4)} items")
for idx, ((wh, req), cand) in enumerate(removed_in_p4.items(), 1):
    print(f"[{idx:2d}] Warehouse: '{wh}'")
    print(f"     Shortage: '{req}'")
    print(f"     Candidate: '{cand}'\n")

print(f"\nChanged candidate in PROC-004: {len(changed_in_p4)} items")
for (wh, req), (old_c, new_c) in changed_in_p4.items():
    print(f"Warehouse: '{wh}' | Shortage: '{req}'")
    print(f"  Old: '{old_c}' -> New: '{new_c}'\n")
