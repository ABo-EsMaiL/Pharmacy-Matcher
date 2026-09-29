import sys
from pathlib import Path
sys.path.insert(0, str(Path(r'd:\AI_Engineer\Pharmacy-agy')))

# Let's inspect each warehouse sheet in results_PROC-001.xlsx
import openpyxl
wb = openpyxl.load_workbook(r'data\processes\PROC-001\results_PROC-001.xlsx')

for sheet in wb.sheetnames:
    if sheet not in ['ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة']:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))[1:]
        print(f"\n=== {sheet} ({len(rows)} matches) ===")
        for i, r in enumerate(rows, 1):
            print(f"  {i:2d}. '{r[0]}'  <--->  '{r[1]}'")
