"""Read the completed new process and check its exported regression outcomes."""
import hashlib
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def main():
    from openpyxl import load_workbook
    from src.fast_match.coordinator import classify_pair, validate_llm_decision

    process_id = sys.argv[1]
    if not process_id.startswith('PROC-') or not process_id[5:].isdigit() or process_id == 'PROC-017':
        raise ValueError('Pass the separate new process ID')
    proc = ROOT / 'data/processes' / process_id
    metadata = json.loads((proc / 'run_metadata.json').read_text(encoding='utf-8'))
    if not all(hashlib.sha256((ROOT / path).read_bytes()).hexdigest() == digest
               for path, digest in metadata['code_hashes'].items()):
        raise ValueError('Historical run used different code; preserve its original validation report')
    results = json.loads((proc / 'results.json').read_text(encoding='utf-8'))
    matches = [(wh, m['shortage_item'], m['warehouse_item'])
               for wh, data in results['warehouse_results'].items() for m in data['matched']]
    reviews = [(wh, r['shortage_item'], r['warehouse_item'])
               for wh, data in results['warehouse_results'].items() for r in data['needs_review']]
    log_path = proc / 'llm_batches.jsonl'
    batches = [json.loads(line) for line in log_path.read_text(encoding='utf-8').splitlines()] if log_path.exists() else []
    adjudicated = set()
    for batch in batches:
        decisions = {d['case_id']: d for d in batch['decisions']}
        for case in batch['inputs']['cases']:
            decision = validate_llm_decision(case, decisions[case['case_id']])
            adjudicated.add((case['requested_name'], decision['warehouse_item'], decision['status']))

    def approved(requested, candidate, status):
        deterministic = classify_pair(requested, candidate)[0]
        return deterministic == status or (deterministic == 'ambiguous' and
                                           (requested, candidate, status) in adjudicated)
    forbidden = [
        ('دولفين 50مجم لبوس 2شريط', 'دولفين 25لبوس 2شريط/باكو90'),
        ('دولفين 12.5مجم لبوس 2شريط س ق', 'دولفين 25لبوس 2شريط/باكو90'),
        ('لبن هيرو بيبي 2', 'هيرو بيبي ١ عادى'),
    ]
    checks = {}
    for requested, candidate in forbidden:
        checks[f'blocked: {requested} -> {candidate}'] = not any(
            r == requested and c == candidate for _, r, c in matches + reviews)
    checks['Flavisef 30ml -> 60ml MATCH'] = any(
        r == 'فلافيسيف شراب30مل' and c == 'فلافيسيف 60 مل شراب باكت 48' for _, r, c in matches)
    checks['all_matches_pass_python'] = all(approved(r, c, 'match') for _, r, c in matches)
    checks['all_reviews_pass_python'] = all(approved(r, c, 'review') for _, r, c in reviews)
    found = {r for _, r, _ in matches}
    review_only = {r for _, r, _ in reviews} - found
    missing = {n['name'] for n in results['not_found']}
    checks['all_unique_shortages_accounted_for'] = (
        len(found | review_only | missing) == results['summary']['unique_shortages']
        and not missing.intersection(found | review_only))

    checks['batch_75_search_off_thinking_high'] = all(
        0 < len(b['inputs']['cases']) <= 75 and
        str(b['gateway']['search_enabled']).lower() == 'false' and
        str(b['gateway']['thinking_level']).lower() == 'high' for b in batches)
    checks['no_rejected_candidate_sent_to_llm'] = all(
        classify_pair(case['requested_name'], c['name'])[0] != 'reject'
        for b in batches for case in b['inputs']['cases'] for c in case['candidates'])
    checks['only_ambiguous_candidates_sent_to_llm'] = all(
        classify_pair(case['requested_name'], c['name'])[0] == 'ambiguous'
        for b in batches for case in b['inputs']['cases'] for c in case['candidates'])
    checks['three_business_outcomes_only'] = (
        'identity_audit' not in results and
        'brand_identity_unverified_count' not in results['summary'])
    checks['source_process_unchanged'] = metadata['source_unchanged']
    checks['run_used_current_code'] = all(
        hashlib.sha256((ROOT / path).read_bytes()).hexdigest() == digest
        for path, digest in metadata['code_hashes'].items())

    book = load_workbook(proc / f'results_{process_id}.xlsx', read_only=True, data_only=False)
    try:
        exported = [(ws.title, row[0], row[1]) for ws in book.worksheets
                    if ws.title not in {'ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة', 'تدقيق هوية البراند'}
                    for row in ws.iter_rows(min_row=2, values_only=True)]
        checks['excel_match_count'] = len(exported) == len(matches)
        checks['excel_missing_count'] = book['لم يُعثر عليه'].max_row - 1 == len(missing)
        checks['excel_review_count'] = book['يحتاج مراجعة'].max_row - 1 == results['summary']['review_count']
        checks['excel_no_identity_audit_sheet'] = 'تدقيق هوية البراند' not in book.sheetnames
        checks['excel_regression_flavisef'] = any(r == 'فلافيسيف شراب30مل' and c == 'فلافيسيف 60 مل شراب باكت 48' for _, r, c in exported)
    finally:
        book.close()
    report = {'process_id': process_id, 'checks': checks, 'passed': all(checks.values()),
              'summary': results['summary'],
              'llm_batches': len(batches), 'unique_matched': len(found),
              'unique_review_only': len(review_only)}
    (proc / 'validation_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if not report['passed']:
        raise SystemExit(1)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
