"""Compare a new run with its offline routing census and observed cache fixtures."""
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'tests'))


def main():
    from test_retrieval_regressions import OBSERVED_PAIRS
    from src.fast_match.search import extract_brand_tokens
    from src.fast_match.attributes import get_strength_pattern
    from src.fast_match.coordinator import validate_llm_decision
    from openpyxl import load_workbook
    process_id = sys.argv[1]
    if not process_id.startswith('PROC-') or not process_id[5:].isdigit() or int(process_id[5:]) <= 19:
        raise ValueError('Only audit a separate process newer than PROC-019')
    folder = ROOT / 'data/processes' / process_id
    census = json.loads((folder / 'retrieval_preflight.json').read_text(encoding='utf-8'))
    batches = [json.loads(line) for line in (folder / 'llm_batches.jsonl').read_text(encoding='utf-8').splitlines()]
    outputs = json.loads((folder / 'results.json').read_text(encoding='utf-8'))
    book = load_workbook(folder / f'results_{process_id}.xlsx', read_only=True)
    try:
        exported_matches = [(row[0], row[1]) for ws in book.worksheets
                            if ws.title not in {'ملخص', 'لم يُعثر عليه', 'يحتاج مراجعة'}
                            for row in ws.iter_rows(min_row=2, values_only=True)]
    finally:
        book.close()
    def identity(name):
        return extract_brand_tokens(name), get_strength_pattern(name)
    fixtures = []
    for query, candidate in OBSERVED_PAIRS:
        seen = []
        for batch in batches:
            decisions = {d['case_id']: d for d in batch['decisions']}
            for case in batch['inputs']['cases']:
                if identity(case['requested_name']) == identity(query) and any(c['name'] == candidate for c in case['candidates']):
                    seen.append({'requested': case['requested_name'], 'decision': validate_llm_decision(case, decisions[case['case_id']])})
        matched = [{'warehouse': wh, **row} for wh, data in outputs['warehouse_results'].items() for row in data['matched']
                   if identity(row['shortage_item']) == identity(query) and row['warehouse_item'] == candidate]
        fixtures.append({'requested_fixture': query, 'candidate': candidate,
                         'reached_llm': bool(seen), 'adjudications': seen, 'match_rows': matched,
                         'excel_match': any(identity(r) == identity(query) and c == candidate
                                            for r, c in exported_matches)})
    actual_cases = sum(len(b['inputs']['cases']) for b in batches)
    actual_pairs = sum(len(c['candidates']) for b in batches for c in b['inputs']['cases'])
    checks = {'actual_cases_equal_preflight': actual_cases == census['totals']['planned_llm_cases'],
              'actual_candidate_pairs_equal_preflight': actual_pairs == census['totals']['planned_llm_candidate_pairs'],
              'seven_observed_candidates_reached_llm': all(f['reached_llm'] for f in fixtures),
              'first_six_observed_pairs_match': all(f['match_rows'] for f in fixtures[:6]),
              'first_six_observed_pairs_present_in_excel': all(f['excel_match'] for f in fixtures[:6]),
              'fixture_json_matches_excel': all(bool(f['match_rows']) == f['excel_match'] for f in fixtures)}
    result = {'process_id': process_id, 'actual_llm_cases': actual_cases,
              'actual_llm_candidate_pairs': actual_pairs, 'llm_batches': len(batches),
              'checks': checks, 'fixtures': fixtures, 'passed': all(checks.values())}
    (folder / 'retrieval_validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result['passed']:
        raise SystemExit(1)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
